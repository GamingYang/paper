import json
import math
import time
from pathlib import Path
from typing import Dict, List, Tuple

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch
import torch.nn as nn
import torch.nn.functional as F
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from torch.utils.data import DataLoader, Dataset


ROOT = Path(r"C:\Users\ROG\Desktop\paper")
CODEX_DIR = ROOT / "codex"
DATA_DIR = ROOT / "PYthon训练代码" / "processed_clean"
OUT_DIR = CODEX_DIR / "paper-update-20260420"
FIG_DIR = OUT_DIR / "figures"

PROGRESS_JSON = OUT_DIR / "paper_update_progress_20260420.json"
PER_BATTERY_XLSX = OUT_DIR / "per_battery_ablation_comparison_20260420.xlsx"
PREDICTIONS_CSV = OUT_DIR / "target_test_predictions_all_models_20260420.csv"
MANUSCRIPT_DOCX = OUT_DIR / "method_experiment_draft_20260420.docx"
RUN_SUMMARY_JSON = OUT_DIR / "paper_update_run_summary_20260420.json"

GLOBAL_MAX_CYCLE = 3000.0
LABEL_MODE = "nominal"
EVAL_BATCH_SIZE = 1024
TIME_LIMIT_SECONDS = 2 * 60 * 60
START_TIME = time.time()

MODELS = [
    {
        "model": "final-retained-v1",
        "use_physical_base": True,
        "use_attention": True,
        "role": "retained complete model",
    },
    {
        "model": "no-physical-base-v1",
        "use_physical_base": False,
        "use_attention": True,
        "role": "ablation without cubic physical degradation base",
    },
    {
        "model": "no-mmd-v1",
        "use_physical_base": True,
        "use_attention": True,
        "role": "ablation without adaptive MMD during training",
    },
    {
        "model": "no-attention-v1",
        "use_physical_base": True,
        "use_attention": False,
        "role": "ablation without temporal attention",
    },
]


def ensure_dirs() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)


def elapsed_seconds() -> float:
    return float(time.time() - START_TIME)


def check_deadline(stage: str) -> None:
    if elapsed_seconds() > TIME_LIMIT_SECONDS:
        write_progress(stage, "timeout", {"elapsed_seconds": elapsed_seconds()})
        raise TimeoutError(f"Time limit exceeded at stage: {stage}")


def write_progress(stage: str, status: str, extra: Dict = None) -> None:
    ensure_dirs()
    payload = {
        "stage": stage,
        "status": status,
        "elapsed_seconds": elapsed_seconds(),
        "time_limit_seconds": TIME_LIMIT_SECONDS,
        "outputs": {
            "out_dir": str(OUT_DIR),
            "figure_dir": str(FIG_DIR),
            "per_battery_xlsx": str(PER_BATTERY_XLSX),
            "predictions_csv": str(PREDICTIONS_CSV),
            "manuscript_docx": str(MANUSCRIPT_DOCX),
            "run_summary_json": str(RUN_SUMMARY_JSON),
        },
    }
    if extra:
        payload.update(extra)
    PROGRESS_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def load_npy(path: Path) -> np.ndarray:
    if not path.exists():
        raise FileNotFoundError(f"Missing required file: {path}")
    return np.load(path)


def load_split(split_name: str, label_mode: str) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    x = load_npy(DATA_DIR / f"clean_{split_name}_x.npy").astype(np.float32)
    p = load_npy(DATA_DIR / f"clean_{split_name}_p.npy").astype(np.float32)
    bidx = load_npy(DATA_DIR / f"clean_{split_name}_battery_index.npy").astype(np.int64)
    batch = load_npy(DATA_DIR / f"clean_{split_name}_batch.npy").astype(np.int64)
    if label_mode == "nominal":
        y = load_npy(DATA_DIR / f"clean_{split_name}_y_nominal.npy").astype(np.float32)
    elif label_mode == "relative":
        y = load_npy(DATA_DIR / f"clean_{split_name}_y_relative.npy").astype(np.float32)
    else:
        raise ValueError("label_mode must be 'nominal' or 'relative'")
    return x, y, p, bidx, batch


class HybridDataset(Dataset):
    def __init__(self, x: np.ndarray, y: np.ndarray, p: np.ndarray, bidx: np.ndarray):
        self.x = x.astype(np.float32)
        self.y = y.astype(np.float32)
        self.p = (p.astype(np.float32) / GLOBAL_MAX_CYCLE).reshape(-1, 1)
        self.bidx = bidx.astype(np.int64)

    def __len__(self) -> int:
        return len(self.x)

    def __getitem__(self, idx: int):
        return self.x[idx], self.y[idx], self.p[idx], self.bidx[idx]


class TemporalAttention(nn.Module):
    def __init__(self, hidden_dim: int):
        super().__init__()
        self.attention = nn.Sequential(
            nn.Linear(hidden_dim, hidden_dim // 2),
            nn.Tanh(),
            nn.Linear(hidden_dim // 2, 1),
        )

    def forward(self, lstm_outputs: torch.Tensor):
        attn_weights = F.softmax(self.attention(lstm_outputs), dim=1)
        context_vector = torch.sum(attn_weights * lstm_outputs, dim=1)
        return context_vector, attn_weights


class ResidualCRNN(nn.Module):
    def __init__(self, base_coeffs: torch.Tensor, use_physical_base: bool, use_attention: bool):
        super().__init__()
        self.base_coeffs = base_coeffs
        self.use_physical_base = use_physical_base
        self.use_attention = use_attention
        self.cnn = nn.Sequential(
            nn.Conv1d(1, 16, 3, padding=1),
            nn.BatchNorm1d(16),
            nn.ReLU(),
            nn.MaxPool1d(2),
            nn.Conv1d(16, 32, 3, padding=1),
            nn.BatchNorm1d(32),
            nn.ReLU(),
            nn.MaxPool1d(2),
        )
        self.lstm = nn.LSTM(32, 64, batch_first=True)
        self.attention = TemporalAttention(hidden_dim=64) if use_attention else None
        self.phy = nn.Sequential(nn.Linear(1, 16), nn.ReLU(), nn.Linear(16, 8), nn.ReLU())
        self.fc_soh = nn.Sequential(nn.Linear(72, 32), nn.ReLU(), nn.Dropout(0.2), nn.Linear(32, 1))

    def forward(self, x: torch.Tensor, p: torch.Tensor):
        if self.use_physical_base:
            a, b, c, d = self.base_coeffs
            base_soh = a * (p**3) + b * (p**2) + c * p + d
        else:
            base_soh = torch.zeros((x.size(0), 1), device=x.device, dtype=x.dtype)

        out = self.cnn(x.unsqueeze(1)).permute(0, 2, 1)
        lstm_out, _ = self.lstm(out)
        if self.use_attention:
            temporal_feat, _ = self.attention(lstm_out)
        else:
            temporal_feat = lstm_out[:, -1, :]
        fused_features = torch.cat((temporal_feat, self.phy(p)), dim=1)
        residual_soh = self.fc_soh(fused_features)
        pred_soh = base_soh + residual_soh if self.use_physical_base else residual_soh
        return pred_soh, fused_features, base_soh


def rmse(y_true: np.ndarray, y_pred: np.ndarray) -> float:
    return float(math.sqrt(np.mean((y_true.reshape(-1) - y_pred.reshape(-1)) ** 2)))


def mae(y_true: np.ndarray, y_pred: np.ndarray) -> float:
    return float(np.mean(np.abs(y_true.reshape(-1) - y_pred.reshape(-1))))


def mape(y_true: np.ndarray, y_pred: np.ndarray) -> float:
    denom = np.clip(np.abs(y_true.reshape(-1)), 1e-8, None)
    return float(np.mean(np.abs((y_true.reshape(-1) - y_pred.reshape(-1)) / denom)) * 100)


def evaluate_predictions(model: nn.Module, loader: DataLoader, device: torch.device) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    model.eval()
    preds_all, trues_all, bidx_all = [], [], []
    with torch.no_grad():
        for x, y, p, b in loader:
            x = x.float().to(device)
            p = p.float().to(device)
            out, _, _ = model(x, p)
            preds_all.append(out.cpu().numpy().reshape(-1))
            trues_all.append(y.numpy().reshape(-1))
            bidx_all.append(b.numpy().reshape(-1))
    return np.concatenate(trues_all), np.concatenate(preds_all), np.concatenate(bidx_all)


def per_battery_metrics(df: pd.DataFrame, pred_col: str) -> pd.DataFrame:
    rows = []
    for battery_index, g in df.groupby("battery_index"):
        rows.append(
            {
                "battery_index": int(battery_index),
                "n_points": int(len(g)),
                "rmse": rmse(g["y_true"].to_numpy(), g[pred_col].to_numpy()),
                "mae": mae(g["y_true"].to_numpy(), g[pred_col].to_numpy()),
                "mape_percent": mape(g["y_true"].to_numpy(), g[pred_col].to_numpy()),
                "cycle_min": float(g["cycle"].min()),
                "cycle_max": float(g["cycle"].max()),
            }
        )
    return pd.DataFrame(rows)


def overall_metrics(df: pd.DataFrame, pred_col: str, model_name: str, best_epoch: int, target_adapt_rmse: float) -> Dict:
    cell_metrics = per_battery_metrics(df, pred_col)
    return {
        "model": model_name,
        "best_epoch_selected_by_target_development": int(best_epoch),
        "target_adapt_best_rmse": float(target_adapt_rmse),
        "target_test_pooled_rmse": rmse(df["y_true"].to_numpy(), df[pred_col].to_numpy()),
        "target_test_macro_rmse": float(cell_metrics["rmse"].mean()),
        "target_test_mae": mae(df["y_true"].to_numpy(), df[pred_col].to_numpy()),
        "target_test_mape_percent": mape(df["y_true"].to_numpy(), df[pred_col].to_numpy()),
    }


def load_model_and_predict(model_cfg: Dict, loader: DataLoader, device: torch.device) -> Tuple[np.ndarray, Dict]:
    model_name = model_cfg["model"]
    checkpoint_path = CODEX_DIR / model_name / f"{model_name}_best_by_target_adapt.pt"
    if not checkpoint_path.exists():
        raise FileNotFoundError(f"Missing checkpoint: {checkpoint_path}")
    checkpoint = torch.load(checkpoint_path, map_location=device, weights_only=False)
    base_coeffs = torch.tensor(checkpoint["base_coeffs"], dtype=torch.float32, device=device)
    model = ResidualCRNN(
        base_coeffs=base_coeffs,
        use_physical_base=model_cfg["use_physical_base"],
        use_attention=model_cfg["use_attention"],
    ).to(device)
    model.load_state_dict(checkpoint["model_state_dict"], strict=True)
    y_true, y_pred, bidx = evaluate_predictions(model, loader, device)
    info = {
        "checkpoint_path": str(checkpoint_path),
        "best_epoch": int(checkpoint.get("epoch", -1)),
        "target_adapt_rmse": float(checkpoint.get("target_adapt_metrics", {}).get("pooled_rmse", np.nan)),
    }
    return y_pred, info


def plot_2x4_blind_test(df: pd.DataFrame, model_name: str, pred_col: str, out_path: Path) -> None:
    batteries = sorted(df["battery_index"].unique().tolist())
    fig, axes = plt.subplots(2, 4, figsize=(18, 8), sharey=True)
    axes = axes.reshape(-1)
    for ax, battery_index in zip(axes, batteries):
        g = df[df["battery_index"] == battery_index].sort_values("cycle")
        cell_rmse = rmse(g["y_true"].to_numpy(), g[pred_col].to_numpy())
        ax.plot(g["cycle"], g["y_true"], color="#222222", linewidth=1.4, label="True SOH")
        ax.plot(g["cycle"], g[pred_col], color="#d95f02", linewidth=1.2, label="Predicted SOH")
        ax.set_title(f"Cell {battery_index} | RMSE={cell_rmse:.4f}", fontsize=10)
        ax.set_xlabel("Cycle")
        ax.grid(True, alpha=0.25)
    axes[0].set_ylabel("SOH")
    axes[4].set_ylabel("SOH")
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="upper center", ncol=2, frameon=False)
    fig.suptitle(f"{model_name}: Batch 10 blind-test SOH trajectories", y=0.98, fontsize=14)
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    fig.savefig(out_path, dpi=300)
    plt.close(fig)


def plot_rmse_bar(rmse_wide: pd.DataFrame, out_path: Path) -> None:
    models = ["final-retained-v1", "no-physical-base-v1", "no-mmd-v1", "no-attention-v1"]
    x = np.arange(len(rmse_wide["battery_index"]))
    width = 0.2
    colors = ["#1b9e77", "#d95f02", "#7570b3", "#e7298a"]
    fig, ax = plt.subplots(figsize=(13, 5.5))
    for idx, model_name in enumerate(models):
        ax.bar(x + (idx - 1.5) * width, rmse_wide[model_name], width=width, label=model_name, color=colors[idx])
    ax.set_xticks(x)
    ax.set_xticklabels(rmse_wide["battery_index"].astype(str).tolist())
    ax.set_xlabel("Battery index in Batch 10")
    ax.set_ylabel("Per-battery RMSE")
    ax.set_title("Per-battery blind-test RMSE comparison")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False, ncol=2)
    fig.tight_layout()
    fig.savefig(out_path, dpi=300)
    plt.close(fig)


def set_run_font(run, font_name="Microsoft YaHei", size=10.5, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(level=level)
    r = p.add_run(text)
    set_run_font(r, size=14 if level == 1 else 12, bold=True)
    return p


def add_para(doc: Document, text: str):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_code_para(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    r = p.add_run(text)
    set_run_font(r, font_name="Consolas", size=9)
    return p


def fill_table(table, rows: List[List]):
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(str(value))
            set_run_font(r, size=9.5, bold=(i == 0))


def write_manuscript_docx(overall_df: pd.DataFrame, rmse_wide: pd.DataFrame, figure_paths: Dict[str, Path]) -> None:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("方法部分与实验设置初稿补充")
    set_run_font(tr, size=16, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("基于严格 Scheme A 协议与 2026-04-20 消融实验结果")
    set_run_font(sr, size=10)

    add_heading(doc, "一、论文表述需要立即调整的事实包", level=1)
    add_para(
        doc,
        "根据本轮严格协议结果，论文叙事应从早期偏重“物理基座与 MMD”的表述，调整为更符合盲测证据的层级结构：物理基座是决定性模块，Attention 是显著有效的辅助时序表征模块，MMD 是小幅但正向的跨域增强模块。该调整并不改变方法总体框架，但会影响创新点排序、消融分析解释和实验结论措辞。",
    )
    add_bullet(doc, "物理基座：负责提供跨电池共享的宏观退化基准线，是当前结果中最稳定、证据最强的组成部分。")
    add_bullet(doc, "残差时序网络：负责学习围绕物理基准线的个体化微观偏差，不应被描述为替代物理模型的黑箱预测器。")
    add_bullet(doc, "Attention：在当前严格协议下去除后 target_test RMSE 明显上升，因此应写作有效的辅助表征模块，但不应提升为核心原创架构。")
    add_bullet(doc, "MMD：去除后 target_test RMSE 仅小幅上升，说明其作用存在但幅度有限，应表述为轻量、克制的跨域正则化增强。")
    add_bullet(doc, "smoothness penalty：当前形式不纳入最终保留模型，论文中不要写成最终有效模块。")
    add_bullet(doc, "Batch 9 不建议只称为 target_adapt，更稳妥的英文表述为 target-domain development set 或 adaptation/development split。")

    add_heading(doc, "二、Methodology 初稿", level=1)
    add_heading(doc, "2.1 Overall Framework", level=2)
    add_para(
        doc,
        "本文面向跨批次锂离子电池 SOH 估计任务，构建一种物理引导的混合迁移学习框架。该框架由全局物理退化基座、残差时序网络和生命周期自适应域对齐约束组成。三者并非简单堆叠，而是承担不同层级的建模职责：物理退化基座描述样本间共享的容量衰减主趋势，残差时序网络刻画不同电池围绕该基准线产生的局部动态偏差，MMD 对齐约束在训练阶段缓解 source batch 与 target-domain development batch 之间的特征分布偏移。",
    )
    add_para(
        doc,
        "设输入片段为 x_i，对应循环位置为 τ_i，SOH 标签为 y_i。模型首先依据归一化循环位置给出物理基准预测，再由时序网络输出残差修正项，最终预测可写为：",
    )
    add_code_para(doc, r"\hat{y}_i = f_{phy}(\tau_i) + f_{res}(x_i, \tau_i)")
    add_para(
        doc,
        "其中 f_{phy}(·) 表示全局物理退化基座，f_{res}(·) 表示残差时序网络。当进行 no-physical-base 消融时，模型退化为仅由残差网络直接输出 SOH 预测。该消融用于检验宏观退化基准线是否真正为跨批次泛化提供有效支撑。",
    )

    add_heading(doc, "2.2 Global Physical Degradation Base", level=2)
    add_para(
        doc,
        "电池容量衰减通常具有相对平滑的长期退化趋势。若完全依赖数据驱动网络从短窗口特征中同时学习宏观趋势和局部波动，模型容易把批次差异、测量噪声和个体差异混入同一表征空间，导致跨批次盲测稳定性下降。因此，本文先采用低阶全局退化基座为所有样本建立共享参考曲线。当前实现中，该基座采用三次多项式形式：",
    )
    add_code_para(doc, r"f_{phy}(\tau)=a\tau^3+b\tau^2+c\tau+d")
    add_para(
        doc,
        "其中 τ 为归一化循环位置，a、b、c、d 为由训练数据估计得到的全局系数。三次多项式在本文中不是被定义为复杂原创结构，而是作为可解释、低复杂度且与退化过程相容的基准函数。其作用是为网络提供宏观退化坐标，使后续神经网络集中学习基准线附近的非线性残差。",
    )

    add_heading(doc, "2.3 Residual Temporal Modeling", level=2)
    add_para(
        doc,
        "在物理基座给出宏观趋势后，不同电池仍会因制造差异、早期活化过程、局部老化速率差异和测量扰动表现出个体化偏差。因此，本文引入残差时序网络对该部分偏差进行建模。残差网络的输入为固定长度的健康相关序列片段及循环位置信息，输出为 SOH 相对于物理基准线的残差修正项。这样，网络学习目标从直接拟合完整 SOH 曲线转变为学习围绕物理基线的局部偏离，有助于降低跨批次迁移难度。",
    )
    add_para(
        doc,
        "具体而言，CNN 模块首先提取局部序列模式，LSTM 模块进一步建模窗口内的时序依赖。Attention 模块用于对不同时间步的隐状态赋予不同权重，从而形成更集中的时序表征。需要强调的是，Attention 在本文中是辅助表征模块，其作用是改善残差特征聚合质量，而不是本文方法的主要原创点。严格协议下的消融结果显示，去除 Attention 后盲测误差明显增加，因此在实验分析中应将其描述为显著有效但功能明确的组成模块。",
    )

    add_heading(doc, "2.4 Adaptive Domain Alignment", level=2)
    add_para(
        doc,
        "由于 source batches 与目标批次在制造批次、循环条件或统计分布上可能存在差异，仅在 source 上进行监督训练可能造成目标批次泛化不足。为缓解这一问题，本文在训练阶段引入生命周期自适应 MMD 对齐约束，使 source 表征与 target-domain development set 的表征在特征空间中保持适度接近。该过程只使用 Batch 9 作为开发/自适应子集，Batch 10 保持为最终 blind-test，不参与模型选择或训练调参。",
    )
    add_code_para(doc, r"\mathcal{L} = \mathcal{L}_{MSE}^{source} + \lambda(\bar{s}_{t}) \mathcal{D}_{MMD}^{2}(z_s,z_t)")
    add_para(
        doc,
        "其中 z_s 与 z_t 分别表示 source 与 target-domain development samples 的中间特征，λ(·) 表示与目标域退化阶段相关的自适应权重。该约束的设计目的不是替代监督误差，而是在训练过程中为跨域特征分布提供轻量正则化。根据当前严格协议结果，MMD 的贡献为正但幅度有限，因此论文中应避免使用“决定性提升”或“根本解决跨域差异”等过度表述。",
    )

    add_heading(doc, "2.5 Optimization and Model Selection", level=2)
    add_para(
        doc,
        "训练过程采用 source supervised learning 与 target-domain development alignment 相结合的策略。前期主要通过 source 标签稳定监督回归，中期加入 MMD 对齐以增强跨批次表征一致性，后期回到纯监督微调以稳定预测误差。最佳 checkpoint 根据 Batch 9 target-domain development set 的 RMSE 选择，Batch 10 仅在 checkpoint 固定后进行一次最终 blind-test。该设置避免了 test-guided checkpoint selection，更符合审稿人对跨域预测任务的实验规范要求。",
    )

    add_heading(doc, "三、Experimental Settings 初稿", level=1)
    add_heading(doc, "3.1 Dataset and Splitting Protocol", level=2)
    add_para(
        doc,
        "实验采用商业 LFP/graphite 锂离子电池寿命数据集，当前预处理版本包含 77 个电池单体。根据固定划分协议，Batch 1-8 被定义为 source domain，用于监督训练；Batch 9 被定义为 target-domain development set，用于目标域自适应和模型选择；Batch 10 被保留为 target-domain blind-test set，仅用于最终泛化评估。该划分方式保证最终测试批次在训练和 checkpoint selection 阶段均不可见。",
    )
    add_para(
        doc,
        "当前数据规模为：source domain 115256 个样本、61 个电池；target-domain development set 16310 个样本、8 个电池；target-domain blind-test set 14556 个样本、8 个电池。输入窗口长度为 100，标签采用 nominal SOH。循环位置被归一化后输入物理基座和位置表征分支。",
    )

    add_heading(doc, "3.2 Compared Models and Ablation Design", level=2)
    add_para(
        doc,
        "为了验证各模块的实际贡献，本文在严格 Scheme A 协议下比较完整保留模型与三个关键消融版本。所有模型使用相同的数据划分、训练轮数和 blind-test 评估方式，以保证对比公平。",
    )
    add_bullet(doc, "final-retained-v1：包含三次物理退化基座、残差时序网络、Attention 和生命周期自适应 MMD。")
    add_bullet(doc, "no-physical-base-v1：移除三次物理退化基座，由残差网络直接预测 SOH，用于检验宏观物理基线的必要性。")
    add_bullet(doc, "no-mmd-v1：关闭训练阶段 MMD 对齐，保留物理基座和 Attention，用于检验跨域对齐约束贡献。")
    add_bullet(doc, "no-attention-v1：移除 Attention，使用 LSTM 最后时刻隐状态作为时序表征，用于检验注意力聚合的辅助作用。")

    add_heading(doc, "3.3 Training Details", level=2)
    add_para(
        doc,
        "所有模型均训练 150 个 epoch，batch size 为 128，初始学习率为 1e-3，并在后期进行学习率衰减。训练过程不保留 smoothness penalty。完整模型和相关消融均采用相同的 checkpoint selection 原则，即以 target-domain development set 上的最优 RMSE 选择模型，再在 Batch 10 上进行一次最终 blind-test。",
    )
    add_para(
        doc,
        "训练记录包括 epoch 级 history.csv、最优 checkpoint、best_metrics.json、summary.xlsx 和 results.xlsx。为了避免测试集泄露，target_test 不参与训练损失、MMD 对齐、早停或模型选择。",
    )

    add_heading(doc, "3.4 Evaluation Metrics", level=2)
    add_para(
        doc,
        "评价指标包括 pooled RMSE、macro RMSE、MAE、MAPE 和 per-battery RMSE。pooled RMSE 反映所有测试样本上的整体误差，macro RMSE 先计算每个电池单体的 RMSE 后再平均，可降低样本数量差异造成的偏置。per-battery RMSE 用于分析性能提升是否集中在特定电池，尤其用于观察 hardest batteries 的误差变化。",
    )
    rows = [["Model", "Best epoch", "Target-dev RMSE", "Blind-test pooled RMSE", "Blind-test macro RMSE"]]
    for _, row in overall_df.iterrows():
        rows.append(
            [
                row["model"],
                int(row["best_epoch_selected_by_target_development"]),
                f"{row['target_adapt_best_rmse']:.6f}",
                f"{row['target_test_pooled_rmse']:.6f}",
                f"{row['target_test_macro_rmse']:.6f}",
            ]
        )
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    fill_table(table, rows)

    add_heading(doc, "3.5 Blind-test Visualization", level=2)
    add_para(
        doc,
        "为补充数值指标，本文进一步绘制 Batch 10 八个电池单体的 2×4 SOH 轨迹图。每个子图给出真实 SOH 与预测 SOH 随循环次数变化的曲线，并在标题中标注对应单体 RMSE。该图用于展示模型在不同目标域电池上的局部误差分布，而不仅依赖单一 pooled RMSE。",
    )
    final_fig = figure_paths.get("final-retained-v1")
    if final_fig and final_fig.exists():
        doc.add_picture(str(final_fig), width=Inches(6.7))
        cap = doc.add_paragraph()
        cr = cap.add_run("Figure: final-retained-v1 on Batch 10 blind-test cells.")
        set_run_font(cr, size=9)

    add_heading(doc, "四、当前严格协议结果的论文解释", level=1)
    add_para(
        doc,
        "当前结果显示，完整模型在 Batch 10 blind-test 上取得 pooled RMSE 0.009499 和 macro RMSE 0.008301。移除物理基座后 pooled RMSE 上升至 0.020716，说明宏观物理退化基线对跨批次泛化具有决定性支撑。移除 Attention 后 pooled RMSE 上升至 0.011677，表明注意力聚合对残差时序表征具有明显辅助作用。移除 MMD 后 pooled RMSE 为 0.009677，仅较完整模型小幅变差，因此 MMD 应被描述为轻量但正向的跨域增强模块。",
    )
    add_para(
        doc,
        "从论文写作角度看，最稳妥的贡献表述应为：本文通过物理退化基座将 SOH 长期趋势从局部残差建模中分离出来，并结合残差时序表征与轻量级目标域开发集对齐，提高了严格 blind-test 场景下的预测稳定性。Attention 和 MMD 均为功能明确的组成模块，但不应被夸大为单独决定性能的原创结构。",
    )

    add_heading(doc, "五、本次运行问题与下一次建议", level=1)
    add_bullet(doc, "本次任务没有重新训练模型，主要耗时来自 checkpoint 推理、绘图、Excel 汇总和 Word 生成，因此远低于 2 小时上限。")
    add_bullet(doc, "第一次运行资产生成脚本时出现 pandas 报错：Per-column arrays must each be 1-dimensional。原因是 y 或 p 等数组保持为二维形状，DataFrame 每一列要求一维数组。已通过 reshape(-1) 修正。下一次生成表格前应先检查数组 shape。")
    add_bullet(doc, "之前出现过后台 Conda DLL 路径问题，本次继续采用显式 Python 解释器前台运行，避免 Start-Process 或后台启动器引发 _ctypes DLL load failed。")
    add_bullet(doc, "文件继续使用英文命名，正文使用中文，避免 PowerShell 中文文件名编码替换问题。")
    add_bullet(doc, "下一次如果要补多 seed，不建议一次性无监控运行全部模型；应先跑 1 epoch smoke test，再分别运行 final-retained 和 no-physical-base 的 seed=7、seed=123。")
    add_bullet(doc, "如果运行超过设定时限，应优先查看 progress JSON、history CSV 和最新 checkpoint，确认已完成阶段后再决定是否续跑。")

    doc.save(MANUSCRIPT_DOCX)


def main():
    ensure_dirs()
    write_progress("start", "running")
    check_deadline("start")

    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    x, y, p, bidx, batch = load_split("target_test", LABEL_MODE)
    dataset = HybridDataset(x, y, p, bidx)
    loader = DataLoader(dataset, batch_size=EVAL_BATCH_SIZE, shuffle=False)
    base_df = pd.DataFrame(
        {
            "sample_index": np.arange(len(y), dtype=np.int64),
            "battery_index": bidx.reshape(-1).astype(int),
            "batch": batch.reshape(-1).astype(int),
            "cycle": p.reshape(-1).astype(float),
            "y_true": y.reshape(-1).astype(float),
        }
    )
    predictions_df = base_df.copy()
    write_progress("data_loaded", "completed", {"device": str(device), "target_test_samples": int(len(y))})

    model_infos = {}
    for model_cfg in MODELS:
        check_deadline(f"predict_{model_cfg['model']}")
        pred, info = load_model_and_predict(model_cfg, loader, device)
        predictions_df[model_cfg["model"]] = pred.astype(float)
        model_infos[model_cfg["model"]] = info
        write_progress(f"predict_{model_cfg['model']}", "completed", {"model_info": info})

    predictions_df.to_csv(PREDICTIONS_CSV, index=False, encoding="utf-8-sig")
    write_progress("predictions_saved", "completed", {"predictions_csv": str(PREDICTIONS_CSV)})

    check_deadline("metrics")
    overall_rows = []
    cell_tables = {}
    for model_cfg in MODELS:
        model_name = model_cfg["model"]
        cell_table = per_battery_metrics(predictions_df, model_name)
        cell_table.insert(0, "model", model_name)
        cell_tables[model_name] = cell_table
        overall_rows.append(
            overall_metrics(
                predictions_df,
                model_name,
                model_name,
                model_infos[model_name]["best_epoch"],
                model_infos[model_name]["target_adapt_rmse"],
            )
        )
    overall_df = pd.DataFrame(overall_rows)

    rmse_wide = pd.DataFrame({"battery_index": sorted(predictions_df["battery_index"].unique().tolist())})
    mae_wide = rmse_wide.copy()
    for model_cfg in MODELS:
        model_name = model_cfg["model"]
        table = cell_tables[model_name][["battery_index", "rmse", "mae", "n_points"]].copy()
        rmse_wide = rmse_wide.merge(table[["battery_index", "rmse"]].rename(columns={"rmse": model_name}), on="battery_index")
        mae_wide = mae_wide.merge(table[["battery_index", "mae"]].rename(columns={"mae": model_name}), on="battery_index")
    n_points = cell_tables["final-retained-v1"][["battery_index", "n_points"]]
    rmse_wide = n_points.merge(rmse_wide, on="battery_index")
    mae_wide = n_points.merge(mae_wide, on="battery_index")

    delta_rows = rmse_wide[["battery_index", "n_points"]].copy()
    for model_name in ["no-physical-base-v1", "no-mmd-v1", "no-attention-v1"]:
        delta_rows[f"{model_name}_minus_final"] = rmse_wide[model_name] - rmse_wide["final-retained-v1"]
    delta_rows["hardest_cell_rank_final"] = rmse_wide["final-retained-v1"].rank(ascending=False, method="min").astype(int)

    with pd.ExcelWriter(PER_BATTERY_XLSX, engine="openpyxl") as writer:
        overall_df.to_excel(writer, sheet_name="overall_metrics", index=False)
        rmse_wide.to_excel(writer, sheet_name="per_battery_rmse", index=False)
        mae_wide.to_excel(writer, sheet_name="per_battery_mae", index=False)
        delta_rows.to_excel(writer, sheet_name="delta_vs_final", index=False)
        for model_name, table in cell_tables.items():
            table.to_excel(writer, sheet_name=model_name[:31], index=False)
    write_progress("excel_saved", "completed", {"per_battery_xlsx": str(PER_BATTERY_XLSX)})

    check_deadline("figures")
    figure_paths = {}
    for model_cfg in MODELS:
        model_name = model_cfg["model"]
        fig_path = FIG_DIR / f"{model_name}_batch10_blind_test_2x4.png"
        plot_2x4_blind_test(predictions_df, model_name, model_name, fig_path)
        figure_paths[model_name] = fig_path
    bar_path = FIG_DIR / "per_battery_rmse_comparison_bar.png"
    plot_rmse_bar(rmse_wide, bar_path)
    write_progress("figures_saved", "completed", {"figures": {k: str(v) for k, v in figure_paths.items()}, "bar": str(bar_path)})

    check_deadline("manuscript")
    write_manuscript_docx(overall_df, rmse_wide, figure_paths)
    write_progress("manuscript_saved", "completed", {"manuscript_docx": str(MANUSCRIPT_DOCX)})

    summary = {
        "status": "success",
        "elapsed_seconds": elapsed_seconds(),
        "time_limit_seconds": TIME_LIMIT_SECONDS,
        "device": str(device),
        "outputs": {
            "out_dir": str(OUT_DIR),
            "per_battery_xlsx": str(PER_BATTERY_XLSX),
            "predictions_csv": str(PREDICTIONS_CSV),
            "manuscript_docx": str(MANUSCRIPT_DOCX),
            "progress_json": str(PROGRESS_JSON),
            "figures": {k: str(v) for k, v in figure_paths.items()},
            "rmse_bar": str(bar_path),
        },
        "overall_metrics": overall_df.to_dict(orient="records"),
        "notes": [
            "No model retraining was performed in this run.",
            "All figures were generated by loading existing best-by-target-adapt checkpoints and evaluating Batch 10 blind-test data.",
            "Target_test remains final blind-test only; no checkpoint selection was performed in this script.",
        ],
    }
    RUN_SUMMARY_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    write_progress("done", "success", {"run_summary_json": str(RUN_SUMMARY_JSON)})
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
