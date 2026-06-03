from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


OUT_DIR = Path(r"C:\Users\ROG\Desktop\paper\codex")
OUT_PATH = OUT_DIR / "codex_run_skill_notes_20260420.docx"


def set_run_font(run, font_name="Microsoft YaHei", size=10.5, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    for line_idx, line in enumerate(text.splitlines()):
        if line_idx:
            p.add_run("\n")
        r = p.add_run(line)
        set_run_font(r, font_name="Consolas", size=9)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        set_run_font(run, size=14 if level == 1 else 12, bold=True)
    if not p.runs:
        r = p.add_run(text)
        set_run_font(r, size=14 if level == 1 else 12, bold=True)
    return p


def fill_table(table, rows):
    for row_idx, row_values in enumerate(rows):
        cells = table.rows[row_idx].cells
        for col_idx, value in enumerate(row_values):
            cells[col_idx].text = ""
            p = cells[col_idx].paragraphs[0]
            r = p.add_run(str(value))
            set_run_font(r, bold=(row_idx == 0))


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Codex 项目运行经验与可复制 Skill 注意事项")
    set_run_font(title_run, size=16, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("适用于锂电池寿命预测论文项目后续训练、文档生成与实验汇报流程")
    set_run_font(sub_run, size=10)

    add_paragraph(doc, "生成日期：2026-04-20")
    add_paragraph(doc, "文档用途：记录本轮项目中出现过的环境、依赖、文档生成、长时间训练和实验记录问题，并提炼为下一次项目可直接复用的操作规范。")

    add_heading(doc, "一、推荐的项目基线", level=1)
    add_bullet(doc, r"项目根目录固定为：C:\Users\ROG\Desktop\paper。")
    add_bullet(doc, r"所有由 Codex 生成的脚本、日志、模型结果和汇报文档优先放入：C:\Users\ROG\Desktop\paper\codex。")
    add_bullet(doc, r"训练环境优先使用明确的解释器路径：D:\Anaconda\envs\battery_project\python.exe。不要依赖“当前激活的是哪个环境”这种隐含状态。")
    add_bullet(doc, r"数据集路径目前为：C:\Users\ROG\Desktop\paper\PYthon训练代码\processed_clean。长任务开始前必须先检查该目录下的 .npy 文件是否存在。")
    add_bullet(doc, "文件命名优先使用英文、数字、短横线和版本号，例如 full-model-v1、no-mmd-v1、ablation-suite-v1。文档正文可以使用中文。")

    add_heading(doc, "二、依赖安装与验证", level=1)
    add_paragraph(doc, "1. Word 文档生成依赖 python-docx。注意：安装包名是 python-docx，但在 Python 中导入时是 docx。")
    add_code(doc, r'''D:\Anaconda\envs\battery_project\python.exe -m pip install python-docx
D:\Anaconda\envs\battery_project\python.exe -c "import docx; print(docx.__version__)"''')
    add_paragraph(doc, "2. Excel 表格生成通常依赖 pandas 和 openpyxl。pandas 负责组织表格数据，openpyxl 负责写入 .xlsx 文件。")
    add_code(doc, r'''D:\Anaconda\envs\battery_project\python.exe -m pip install pandas openpyxl
D:\Anaconda\envs\battery_project\python.exe -c "import pandas, openpyxl; print('excel-ok')"''')
    add_paragraph(doc, "3. 如果你在 base 环境安装了包，但 Codex 调用的是 battery_project 环境，base 中安装成功并不代表项目环境可用。以后必须把包安装到实际运行训练代码的环境中。")
    add_paragraph(doc, "4. 对本项目而言，已经验证 battery_project 环境可以调用 python-docx、pandas、openpyxl、torch，并可以创建 Word、Excel、JSON、CSV 和模型 checkpoint 文件。")

    add_heading(doc, "三、ripgrep rg 安装经验", level=1)
    add_paragraph(doc, "rg 用于快速搜索文件和文本，比传统 grep 或 PowerShell 全目录扫描更适合大型项目。")
    add_code(doc, r"""choco install ripgrep -y
rg --version""")
    add_paragraph(doc, "本轮遇到的问题包括：Chocolatey 社区源返回 503 Service Unavailable、提示 Unable to find package 'ripgrep'、安装后 rg 不是内部或外部命令。")
    add_bullet(doc, "503 通常是 Chocolatey 社区源临时不可用或网络访问问题，不一定是命令写错。处理方式是稍后重试，或检查网络/代理/Chocolatey source。")
    add_bullet(doc, r"如果安装成功但 rg 无法识别，先关闭并重新打开 PowerShell，再执行 rg --version。若仍失败，检查 C:\ProgramData\chocolatey\bin 是否在 PATH 中。")
    add_bullet(doc, "验证标准是 rg --version 能输出 ripgrep 版本号。本轮最终验证版本为 ripgrep 14.1.0。")

    add_heading(doc, "四、Word 与 Excel 文档生成规则", level=1)
    add_bullet(doc, "Word：使用 python-docx 生成 .docx，适合写论文初稿、实验汇报、运行总结。公式若未确认复杂排版，优先以 LaTeX 代码文本形式写入。")
    add_bullet(doc, "Excel：使用 pandas + openpyxl 生成 .xlsx，适合保存每轮训练指标、单体电池结果、总 RMSE 汇总表。")
    add_bullet(doc, "JSON：用于保存机器可读的完整配置和指标，后续可以自动读取并生成汇总文档。")
    add_bullet(doc, "CSV：用于保存 epoch 级训练历史，便于后续画曲线或排查训练震荡。")
    add_bullet(doc, "不要把所有结果散放在项目根目录。每次实验都应建立独立子目录，例如 codex/full-model-v1、codex/no-mmd-v1。")
    add_bullet(doc, "PowerShell 管道中传递中文文件名曾出现编码替换问题。后续建议：文件名用英文，正文用中文；涉及路径时使用 Python pathlib 或 PowerShell -LiteralPath。")

    add_heading(doc, "五、长时间训练与自动停止策略", level=1)
    add_paragraph(doc, "本轮曾出现训练运行 3 个多小时但没有得到有效结果的问题。后续长任务必须具备三层保护：短任务验证、前台直接运行、超时控制。")
    add_bullet(doc, "短任务验证：正式训练前先跑 1 个 epoch smoke test，确认数据路径、模型前向传播、loss、Excel/JSON 写入都正常。")
    add_bullet(doc, "语法验证：长任务前先运行 py_compile，避免训练启动后才暴露缩进、拼写或导入错误。")
    add_code(doc, r"""D:\Anaconda\envs\battery_project\python.exe -m py_compile C:\Users\ROG\Desktop\paper\codex\your_script.py""")
    add_bullet(doc, "前台直接运行：优先使用明确 Python 路径直接运行脚本，不优先使用后台启动器。")
    add_code(doc, r"""D:\Anaconda\envs\battery_project\python.exe C:\Users\ROG\Desktop\paper\codex\your_script.py""")
    add_bullet(doc, "Codex 工具层运行长任务时设置 timeout_ms。例如 1 小时对应 3600000 ms。超过该时间工具会终止运行，避免无限等待。")
    add_bullet(doc, "训练脚本应每个 epoch 打印关键指标，并持续写入 history.csv 或 JSON。即使中途失败，也能保留部分诊断信息。")

    add_heading(doc, "六、Conda DLL 与后台启动问题", level=1)
    add_paragraph(doc, "本轮出现过 pandas -> ctypes -> _ctypes DLL load failed。直接调用 battery_project/python.exe 时正常，但后台启动器中失败，说明问题主要来自 Conda 环境的 DLL 搜索路径没有补齐。")
    add_paragraph(doc, "如果未来确实需要后台运行，启动前应显式补齐 CONDA_PREFIX、Library\\bin、DLLs、Scripts 等路径。但从本项目实践看，优先推荐前台直接运行并设置超时。")
    add_code(doc, r'''$env:CONDA_PREFIX="D:\Anaconda\envs\battery_project"
$env:PATH="$env:CONDA_PREFIX;$env:CONDA_PREFIX\Library\bin;$env:CONDA_PREFIX\DLLs;$env:CONDA_PREFIX\Scripts;$env:PATH"
D:\Anaconda\envs\battery_project\python.exe -c "import ctypes, pandas; print('dll-ok')"''')
    add_bullet(doc, "Windows 环境变量名大小写不敏感。曾出现 Path 与 PATH 重复加入导致的异常：Item has already been added. Key in dictionary: 'Path' Key being added: 'PATH'。后续不要在同一个环境变量字典中同时设置 Path 和 PATH。")
    add_bullet(doc, "如果脚本在后台很快退出，而前台直接运行正常，优先怀疑后台启动器环境变量不完整，而不是模型代码本身错误。")

    add_heading(doc, "七、实验输出目录与记录规范", level=1)
    add_paragraph(doc, "以后每个训练任务建议采用以下目录结构。")
    add_code(doc, r"""codex/
  full-model-v1/
    train_full_model.py
    full-model-v1_best_by_target_adapt.pt
    full-model-v1_best_metrics.json
    full-model-v1_history.csv
    full-model-v1_results.xlsx
    full-model-v1_summary.xlsx
  no-mmd-v1/
  no-attention-v1/
  ablation-suite-v1/
    ablation-suite-v1_report.docx
  model_runs_summary.xlsx""")
    add_bullet(doc, "每个模型子文件夹至少保存：训练脚本、最佳 checkpoint、best_metrics.json、history.csv、results.xlsx、summary.xlsx。")
    add_bullet(doc, "全局汇总表 model_runs_summary.xlsx 只保留关键列：model、best_target_adapt_rmse、final_target_test_rmse。")
    add_bullet(doc, "JSON 中建议记录：Python 路径、device、seed、epoch、batch_size、学习率、数据路径、数据划分协议、best_epoch、target_adapt RMSE、target_test RMSE。")
    add_bullet(doc, "训练完成后再生成 Word 汇报文档，避免手工复制数字导致误差。")

    add_heading(doc, "八、论文实验的科学性约束", level=1)
    add_bullet(doc, "必须保持严格 Scheme A：source 使用 Batch 1-8，target_adapt 使用 Batch 9，target_test 使用 Batch 10。")
    add_bullet(doc, "target_test 只能在最佳 checkpoint 确定后进行最终 blind-test，不能参与 epoch 选择、调参或早停判断。")
    add_bullet(doc, "物理基座、残差网络、Attention、MMD 的表述要克制。物理基座负责宏观趋势，残差网络负责个体化微观波动，Attention 是辅助表征模块，MMD 是训练阶段的跨批次分布对齐约束。")
    add_bullet(doc, "当前 smoothness penalty 不作为最终保留模块。论文中不要把它写成最终有效贡献。")
    add_bullet(doc, "MMD 的贡献需要根据实际消融结果谨慎表述。本轮 no-MMD 与完整模型差异较小，因此不宜写成决定性创新。")
    add_bullet(doc, "如果要增强 SCI 审稿说服力，后续建议对完整模型和关键消融做多随机种子重复实验，并报告均值和标准差。")

    add_heading(doc, "九、常见问题速查表", level=1)
    rows = [
        ["问题现象", "可能原因", "推荐处理"],
        ["No module named docx", "python-docx 未安装到当前运行环境", r"用 D:\Anaconda\envs\battery_project\python.exe -m pip install python-docx 安装并验证 import docx"],
        ["Excel 无法写入 .xlsx", "缺少 openpyxl 或 pandas", "安装 pandas openpyxl，并用最小表格写入测试"],
        ["rg 不是内部或外部命令", "ripgrep 未装成功或 PATH 未刷新", "重新打开 PowerShell，执行 rg --version；必要时检查 Chocolatey bin 路径"],
        ["Chocolatey 503", "社区源临时不可用或网络问题", "稍后重试，检查 choco source list，不要误判为命令错误"],
        ["中文文件名传参乱码", "PowerShell 管道编码替换", "文件名用英文；正文用中文；路径用 pathlib 或 -LiteralPath"],
        ["_ctypes DLL load failed", "后台进程未加载 Conda DLL 路径", "优先前台直接运行；必要时补齐 CONDA_PREFIX、Library\\bin、DLLs、Scripts"],
        ["Path/PATH 重复键错误", "Windows 环境变量大小写不敏感", "不要同时构造 Path 和 PATH 两个键"],
        ["长时间无结果", "无超时、无日志、后台进程异常", "先 1 epoch smoke test，再 py_compile，正式运行设置 timeout_ms"],
        ["RMSE 结果不可信", "target_test 泄露或调参使用测试集", "坚持 target_adapt 选模型，target_test 只最终 blind-test 一次"],
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    fill_table(table, rows)

    add_heading(doc, "十、下一次项目可直接复制的 Skill 提示词", level=1)
    add_paragraph(doc, "以下内容可以直接复制到下一次项目的前置说明或任务开头，用于约束 Codex 的运行方式。")
    add_code(doc, r"""请使用以下固定规则推进本项目：
1. 项目根目录为 C:\Users\ROG\Desktop\paper，所有新生成文件放入 codex 子目录。
2. Python 解释器固定使用 D:\Anaconda\envs\battery_project\python.exe，不依赖 base 环境或隐式激活环境。
3. 长任务前必须先做 py_compile 和 1 epoch smoke test。
4. 正式训练必须设置超时；30 分钟任务用 1800000 ms，1 小时任务用 3600000 ms。
5. 每个模型单独建立英文命名子目录，例如 full-model-v1、no-mmd-v1。
6. 每次训练至少输出 checkpoint、best_metrics.json、history.csv、results.xlsx、summary.xlsx。
7. 全局汇总表只记录 model、best_target_adapt_rmse、final_target_test_rmse。
8. Word 文档正文可用中文，但文件名优先使用英文，避免 PowerShell 中文文件名编码问题。
9. 训练协议必须保持严格 Scheme A：source=Batch 1-8，target_adapt=Batch 9，target_test=Batch 10 且只最终盲测一次。
10. 论文表述保持克制，不能把 Attention、MMD 或未保留的 smoothness 写成核心原创贡献。""")

    add_heading(doc, "十一、本轮已验证能力", level=1)
    add_bullet(doc, "可以在 battery_project 环境中调用 Python 脚本并记录运行结果。")
    add_bullet(doc, "可以生成 Word 文档，包括论文初稿、实验汇报和本经验总结。")
    add_bullet(doc, "可以生成 Excel 表格，包括每轮训练结果和总 RMSE 汇总表。")
    add_bullet(doc, "可以生成 JSON、CSV 和 PyTorch checkpoint，用于复现实验记录。")
    add_bullet(doc, "可以运行完整 150 epoch 训练和三个消融实验，并通过工具层超时控制避免无限等待。")

    add_paragraph(doc, "结论：下一次项目应把“固定解释器、英文文件名、先小跑后长跑、前台直接运行、超时控制、分目录保存、严格数据划分、防止测试集泄露”作为默认执行规范。")

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
