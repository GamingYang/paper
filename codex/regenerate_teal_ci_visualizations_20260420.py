import importlib.util
import json
import math
import time
from pathlib import Path
from typing import Dict, List, Tuple

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from torch.utils.data import DataLoader


ROOT = Path(r"C:\Users\ROG\Desktop\paper")
CODEX_DIR = ROOT / "codex"
BASE_SCRIPT = CODEX_DIR / "generate_paper_update_assets_20260420.py"
OUT_DIR = CODEX_DIR / "visualization-update-20260420"
FIG_DIR = OUT_DIR / "figures"
REPORT_DOCX = OUT_DIR / "teal_ci_visualization_report_20260420.docx"
SUMMARY_XLSX = OUT_DIR / "teal_ci_visualization_summary_20260420.xlsx"
RUN_JSON = OUT_DIR / "teal_ci_visualization_run_summary_20260420.json"
PROGRESS_JSON = OUT_DIR / "teal_ci_visualization_progress_20260420.json"

TIME_LIMIT_SECONDS = 30 * 60
START_TIME = time.time()
EVAL_BATCH_SIZE = 1024

COLORS = {
    "true": "#163B3A",
    "pred": "#00A7A5",
    "band": "#84DCCF",
    "error": "#007C89",
    "zero": "#5B7F7A",
    "grid": "#B7D8D2",
}


def import_base_module():
    spec = importlib.util.spec_from_file_location("paper_assets_20260420", BASE_SCRIPT)
    if spec is None or spec.loader is None:
        raise ImportError(f"Cannot import base script: {BASE_SCRIPT}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


assets = import_base_module()


def ensure_dirs() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)


def elapsed_seconds() -> float:
    return float(time.time() - START_TIME)


def write_progress(stage: str, status: str, extra: Dict = None) -> None:
    ensure_dirs()
    payload = {
        "stage": stage,
        "status": status,
        "elapsed_seconds": elapsed_seconds(),
        "time_limit_seconds": TIME_LIMIT_SECONDS,
        "outputs": {
            "out_dir": str(OUT_DIR),
            "figure_dir": str(FIG_DIR),
            "report_docx": str(REPORT_DOCX),
            "summary_xlsx": str(SUMMARY_XLSX),
            "run_json": str(RUN_JSON),
        },
    }
    if extra:
        payload.update(extra)
    PROGRESS_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def check_deadline(stage: str) -> None:
    if elapsed_seconds() > TIME_LIMIT_SECONDS:
        write_progress(stage, "timeout")
        raise TimeoutError(f"Time limit exceeded at stage: {stage}")


def make_loader(split_name: str):
    x, y, p, bidx, batch = assets.load_split(split_name, assets.LABEL_MODE)
    dataset = assets.HybridDataset(x, y, p, bidx)
    loader = DataLoader(dataset, batch_size=EVAL_BATCH_SIZE, shuffle=False)
    meta = pd.DataFrame(
        {
            "sample_index": np.arange(len(y), dtype=np.int64),
            "battery_index": bidx.reshape(-1).astype(int),
            "batch": batch.reshape(-1).astype(int),
            "cycle": p.reshape(-1).astype(float),
            "y_true": y.reshape(-1).astype(float),
        }
    )
    return loader, meta


def load_checkpoint_model(model_cfg: Dict, device: torch.device):
    model_name = model_cfg["model"]
    checkpoint_path = CODEX_DIR / model_name / f"{model_name}_best_by_target_adapt.pt"
    if not checkpoint_path.exists():
        raise FileNotFoundError(f"Missing checkpoint: {checkpoint_path}")
    checkpoint = torch.load(checkpoint_path, map_location=device, weights_only=False)
    base_coeffs = torch.tensor(checkpoint["base_coeffs"], dtype=torch.float32, device=device)
    model = assets.ResidualCRNN(
        base_coeffs=base_coeffs,
        use_physical_base=model_cfg["use_physical_base"],
        use_attention=model_cfg["use_attention"],
    ).to(device)
    model.load_state_dict(checkpoint["model_state_dict"], strict=True)
    model.eval()
    info = {
        "checkpoint_path": str(checkpoint_path),
        "best_epoch": int(checkpoint.get("epoch", -1)),
        "target_adapt_checkpoint_rmse": float(checkpoint.get("target_adapt_metrics", {}).get("pooled_rmse", np.nan)),
    }
    return model, info


def predict(model, loader, device: torch.device) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    return assets.evaluate_predictions(model, loader, device)


def residual_band_from_development(y_true: np.ndarray, y_pred: np.ndarray) -> Dict:
    residual = y_pred.reshape(-1) - y_true.reshape(-1)
    std = float(np.std(residual, ddof=1))
    mean = float(np.mean(residual))
    half_width = float(1.96 * std)
    return {
        "development_residual_mean": mean,
        "development_residual_std": std,
        "ci_half_width_95": half_width,
        "ci_note": "Prediction band estimated as pred +/- 1.96 * std(residuals on Batch 9 target-domain development set).",
    }


def cell_rmse(y_true: np.ndarray, y_pred: np.ndarray) -> float:
    return float(math.sqrt(np.mean((y_true.reshape(-1) - y_pred.reshape(-1)) ** 2)))


def plot_enhanced_2x4(df: pd.DataFrame, model_name: str, ci_half_width: float, out_path: Path) -> Dict:
    batteries = sorted(df["battery_index"].unique().tolist())
    y_low = min(float(df["y_true"].min()), float(df[model_name].min() - ci_half_width))
    y_high = max(float(df["y_true"].max()), float(df[model_name].max() + ci_half_width))
    y_pad = 0.04 * max(y_high - y_low, 1e-6)
    y_lim = (y_low - y_pad, y_high + y_pad)

    fig, axes = plt.subplots(2, 4, figsize=(18.5, 8.8), sharey=True)
    axes = axes.reshape(-1)
    cell_rows = []

    for idx, (ax, battery_index) in enumerate(zip(axes, batteries)):
        g = df[df["battery_index"] == battery_index].sort_values("cycle")
        cycle = g["cycle"].to_numpy()
        y_true = g["y_true"].to_numpy()
        y_pred = g[model_name].to_numpy()
        error = y_pred - y_true
        rmse_val = cell_rmse(y_true, y_pred)
        mae_val = float(np.mean(np.abs(error)))
        err_lim = max(float(np.max(np.abs(error))) * 1.25, ci_half_width * 0.35, 0.002)

        band = ax.fill_between(
            cycle,
            y_pred - ci_half_width,
            y_pred + ci_half_width,
            color=COLORS["band"],
            alpha=0.28,
            linewidth=0,
            label="95% band",
        )
        true_line = ax.plot(cycle, y_true, color=COLORS["true"], linewidth=1.45, label="True SOH")[0]
        pred_line = ax.plot(cycle, y_pred, color=COLORS["pred"], linewidth=1.35, label="Predicted SOH")[0]

        ax2 = ax.twinx()
        zero_line = ax2.axhline(0, color=COLORS["zero"], linewidth=0.75, alpha=0.65)
        err_line = ax2.plot(cycle, error, color=COLORS["error"], linewidth=0.85, linestyle="--", alpha=0.9, label="Error")[0]
        ax2.set_ylim(-err_lim, err_lim)
        ax2.tick_params(axis="y", labelsize=7, colors=COLORS["error"])
        if idx % 4 == 3:
            ax2.set_ylabel("Error (Pred-True)", color=COLORS["error"], fontsize=8)
        else:
            ax2.set_yticklabels([])

        ax.set_ylim(*y_lim)
        ax.set_title(f"Cell {battery_index} | RMSE={rmse_val:.4f}", fontsize=9.5, loc="left", pad=4)
        ax.set_xlabel("Cycle", fontsize=8)
        if idx % 4 == 0:
            ax.set_ylabel("SOH", fontsize=8)
        ax.grid(True, color=COLORS["grid"], alpha=0.38, linewidth=0.6)
        ax.tick_params(axis="both", labelsize=7)

        handles = [true_line, pred_line, band, err_line]
        labels = ["True", "Pred", "95% band", "Error"]
        ax.legend(handles, labels, loc="upper right", fontsize=6.2, frameon=True, framealpha=0.86, borderpad=0.35)

        cell_rows.append(
            {
                "model": model_name,
                "battery_index": int(battery_index),
                "rmse": rmse_val,
                "mae": mae_val,
                "max_abs_error": float(np.max(np.abs(error))),
                "n_points": int(len(g)),
            }
        )

    fig.suptitle(
        f"{model_name}: Batch 10 blind-test SOH prediction with development-set 95% residual band",
        fontsize=14,
        color=COLORS["true"],
        y=0.985,
    )
    fig.tight_layout(rect=[0, 0, 1, 0.955], w_pad=1.6, h_pad=1.35)
    fig.savefig(out_path, dpi=300)
    plt.close(fig)
    return {"figure": str(out_path), "cell_metrics": cell_rows}


def set_run_font(run, font_name="Microsoft YaHei", size=10.5, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(level=level)
    r = p.add_run(text)
    set_run_font(r, size=14 if level == 1 else 12, bold=True)


def add_para(doc: Document, text: str):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r)


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run_font(r)


def fill_table(table, rows: List[List]):
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(str(value))
            set_run_font(r, size=9, bold=(i == 0))


def write_report(overall_df: pd.DataFrame, figure_paths: Dict[str, str], issues: List[str]) -> None:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("青绿色系 2×4 盲测可视化更新汇报")
    set_run_font(tr, size=16, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("图例位置、95% 经验误差带与右轴误差曲线修订")
    set_run_font(sr, size=10)

    add_heading(doc, "一、修改目标", level=1)
    add_bullet(doc, "将图例从全局顶部移动到每个电池子图右上角，避免与主标题或子图标题重叠。")
    add_bullet(doc, "为预测曲线增加 95% 经验误差带。该误差带由 Batch 9 target-domain development set 的残差标准差估计，不使用 Batch 10 盲测误差反推区间。")
    add_bullet(doc, "在每个子图中增加右侧纵坐标，绘制 signed error = Predicted SOH - True SOH，并加入零误差参考线。")
    add_bullet(doc, "整体配色改为青绿色系，保持真实值、预测值、误差带和误差曲线之间的视觉区分。")

    add_heading(doc, "二、区间定义说明", level=1)
    add_para(
        doc,
        "当前模型为确定性回归模型，未训练贝叶斯网络、深度集成或分位数回归模型，因此严格意义上的预测置信区间无法直接由模型输出。为避免夸大不确定性建模能力，本次图中使用的是基于 Batch 9 target-domain development residual 的经验 95% 误差带，定义为 predicted SOH +/- 1.96 x std(residual_dev)。该带状区域用于辅助观察误差尺度，不应在论文中表述为严格概率校准后的置信区间。",
    )

    add_heading(doc, "三、图形参数与关键指标", level=1)
    rows = [["Model", "Best epoch", "Target-dev RMSE", "Batch 10 RMSE", "Dev residual std", "95% half-width"]]
    for _, row in overall_df.iterrows():
        rows.append(
            [
                row["model"],
                int(row["best_epoch"]),
                f"{row['target_adapt_rmse']:.6f}",
                f"{row['target_test_rmse']:.6f}",
                f"{row['development_residual_std']:.6f}",
                f"{row['ci_half_width_95']:.6f}",
            ]
        )
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    fill_table(table, rows)

    add_heading(doc, "四、生成图像", level=1)
    for model_name, fig_path in figure_paths.items():
        add_para(doc, f"{model_name}")
        doc.add_picture(fig_path, width=Inches(6.7))

    add_heading(doc, "五、本次运行问题与处理", level=1)
    if issues:
        for item in issues:
            add_bullet(doc, item)
    else:
        add_bullet(doc, "本次重新绘图未发生新的运行错误。")
    add_bullet(doc, "本次没有重新训练模型，只加载已有 best-by-target-adapt checkpoint 做 Batch 9/Batch 10 推理和绘图，因此运行时间显著低于半小时上限。")
    add_bullet(doc, "下一次如果继续要求置信区间，建议明确采用经验残差带、MC Dropout、深度集成或分位数回归中的哪一种；如果不额外训练模型，最稳妥表述仍是“经验误差带”。")

    doc.save(REPORT_DOCX)


def main():
    ensure_dirs()
    issues = []
    write_progress("start", "running")
    check_deadline("start")

    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    adapt_loader, _ = make_loader("target_adapt")
    test_loader, test_meta = make_loader("target_test")
    write_progress("data_loaded", "completed", {"device": str(device)})

    prediction_df = test_meta.copy()
    overall_rows = []
    figure_paths = {}
    all_cell_rows = []

    for model_cfg in assets.MODELS:
        check_deadline(f"predict_and_plot_{model_cfg['model']}")
        model_name = model_cfg["model"]
        model, ckpt_info = load_checkpoint_model(model_cfg, device)

        y_adapt, pred_adapt, _ = predict(model, adapt_loader, device)
        band_info = residual_band_from_development(y_adapt, pred_adapt)

        y_test, pred_test, bidx_test = predict(model, test_loader, device)
        prediction_df[model_name] = pred_test.astype(float)
        test_rmse = assets.rmse(y_test, pred_test)
        test_mae = assets.mae(y_test, pred_test)
        test_mape = assets.mape(y_test, pred_test)

        fig_path = FIG_DIR / f"{model_name}_batch10_blind_test_2x4_teal_ci_error.png"
        plot_result = plot_enhanced_2x4(prediction_df, model_name, band_info["ci_half_width_95"], fig_path)
        figure_paths[model_name] = str(fig_path)
        all_cell_rows.extend(plot_result["cell_metrics"])

        overall_rows.append(
            {
                "model": model_name,
                "best_epoch": ckpt_info["best_epoch"],
                "target_adapt_rmse": ckpt_info["target_adapt_checkpoint_rmse"],
                "target_test_rmse": test_rmse,
                "target_test_mae": test_mae,
                "target_test_mape_percent": test_mape,
                **band_info,
                "figure_path": str(fig_path),
                "checkpoint_path": ckpt_info["checkpoint_path"],
            }
        )
        write_progress(f"figure_saved_{model_name}", "completed", {"figure": str(fig_path)})

    overall_df = pd.DataFrame(overall_rows)
    cell_df = pd.DataFrame(all_cell_rows)
    with pd.ExcelWriter(SUMMARY_XLSX, engine="openpyxl") as writer:
        overall_df.to_excel(writer, sheet_name="overall", index=False)
        cell_df.to_excel(writer, sheet_name="cell_metrics", index=False)
        prediction_df.to_excel(writer, sheet_name="target_test_predictions", index=False)
    write_progress("summary_xlsx_saved", "completed", {"summary_xlsx": str(SUMMARY_XLSX)})

    write_report(overall_df, figure_paths, issues)
    write_progress("report_saved", "completed", {"report_docx": str(REPORT_DOCX)})

    summary = {
        "status": "success",
        "elapsed_seconds": elapsed_seconds(),
        "time_limit_seconds": TIME_LIMIT_SECONDS,
        "device": str(device),
        "outputs": {
            "out_dir": str(OUT_DIR),
            "figures": figure_paths,
            "summary_xlsx": str(SUMMARY_XLSX),
            "report_docx": str(REPORT_DOCX),
            "progress_json": str(PROGRESS_JSON),
        },
        "overall": overall_df.to_dict(orient="records"),
        "issues": issues,
    }
    RUN_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    write_progress("done", "success", {"run_json": str(RUN_JSON)})
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
