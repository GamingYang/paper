import json
import math
import sys
from pathlib import Path
from typing import Dict, List, Tuple

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
import torch.nn.functional as F
import torch.optim as optim
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from torch.utils.data import DataLoader, Dataset


ROOT = Path(r"C:\Users\ROG\Desktop\paper")
CODEX_DIR = ROOT / "codex"
DATA_DIR = ROOT / "PYthon训练代码" / "processed_clean"
SUITE_DIR = CODEX_DIR / "ablation-suite-v1"
SUMMARY_XLSX = CODEX_DIR / "model_runs_summary.xlsx"
REPORT_DOCX = SUITE_DIR / "ablation-suite-v1_report.docx"

LABEL_MODE = "nominal"
DATA_PROTOCOL = "processed_clean_fixed_split"
SEED = 42
EPOCHS = 150
BATCH_SIZE = 128
EVAL_BATCH_SIZE = 1024
LEARNING_RATE = 1e-3
GLOBAL_MAX_CYCLE = 3000.0
TARGET_MMD_WEIGHT = 0.00005
LAMBDA_SMOOTH = 0.0
LR_DECAY_EPOCH = 100
LR_DECAY_FACTOR = 0.1

VARIANTS = [
    {
        "model": "no-physical-base-v1",
        "use_physical_base": False,
        "use_attention": True,
        "use_mmd": True,
        "description": "移除三次物理退化基座，保留残差时序网络、Attention 与 MMD。",
    },
    {
        "model": "no-mmd-v1",
        "use_physical_base": True,
        "use_attention": True,
        "use_mmd": False,
        "description": "关闭生命周期自适应 MMD，保留物理基座、残差时序网络与 Attention。",
    },
    {
        "model": "no-attention-v1",
        "use_physical_base": True,
        "use_attention": False,
        "use_mmd": True,
        "description": "移除 Attention，使用 LSTM 最后时刻隐藏状态作为时序表征，保留物理基座与 MMD。",
    },
]


def set_seed(seed: int) -> None:
    np.random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(seed)
    if torch.backends.cudnn.is_available():
        torch.backends.cudnn.deterministic = True
        torch.backends.cudnn.benchmark = False


def load_npy(path: Path) -> np.ndarray:
    if not path.exists():
        raise FileNotFoundError(f"Missing required file: {path}")
    return np.load(path)


def load_split(split_name: str, label_mode: str) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    x = load_npy(DATA_DIR / f"clean_{split_name}_x.npy").astype(np.float32)
    p = load_npy(DATA_DIR / f"clean_{split_name}_p.npy").astype(np.float32)
    bidx = load_npy(DATA_DIR / f"clean_{split_name}_battery_index.npy").astype(np.int64)
    batch = load_npy(DATA_DIR / f"clean_{split_name}_batch.npy").astype(np.int64)
    if label_mode == "nominal":
        y = load_npy(DATA_DIR / f"clean_{split_name}_y_nominal.npy").astype(np.float32)
    elif label_mode == "relative":
        y = load_npy(DATA_DIR / f"clean_{split_name}_y_relative.npy").astype(np.float32)
    else:
        raise ValueError("label_mode must be 'nominal' or 'relative'")
    return x, y, p, bidx, batch


class HybridDataset(Dataset):
    def __init__(self, x: np.ndarray, y: np.ndarray, p: np.ndarray, bidx: np.ndarray):
        self.x = x.astype(np.float32)
        self.y = y.astype(np.float32)
        self.p = (p.astype(np.float32) / GLOBAL_MAX_CYCLE).reshape(-1, 1)
        self.bidx = bidx.astype(np.int64)

    def __len__(self) -> int:
        return len(self.x)

    def __getitem__(self, idx: int):
        return self.x[idx], self.y[idx], self.p[idx], self.bidx[idx]


class LifetimeAdaptiveMMD(nn.Module):
    def __init__(self, kernel_mul: float = 2.0, kernel_num: int = 5, lambda_scale: float = 5.5):
        super().__init__()
        self.kernel_num = kernel_num
        self.kernel_mul = kernel_mul
        self.lambda_scale = lambda_scale

    def gaussian_kernel(self, source: torch.Tensor, target: torch.Tensor) -> torch.Tensor:
        n_samples = int(source.size(0)) + int(target.size(0))
        total = torch.cat([source, target], dim=0)
        total0 = total.unsqueeze(0).expand(total.size(0), total.size(0), total.size(1))
        total1 = total.unsqueeze(1).expand(total.size(0), total.size(0), total.size(1))
        l2_distance = ((total0 - total1) ** 2).sum(2)
        bandwidth = torch.sum(l2_distance.detach()) / max(n_samples**2 - n_samples, 1)
        bandwidth = torch.clamp(bandwidth, min=1e-6)
        bandwidth = bandwidth / (self.kernel_mul ** (self.kernel_num // 2))
        bandwidth_list = [bandwidth * (self.kernel_mul**i) for i in range(self.kernel_num)]
        return sum(torch.exp(-l2_distance / bw) for bw in bandwidth_list)

    def compute_mmd(self, source: torch.Tensor, target: torch.Tensor) -> torch.Tensor:
        if source.size(0) == 0 or target.size(0) == 0:
            return torch.tensor(0.0, device=source.device)
        batch_size_s = int(source.size(0))
        kernels = self.gaussian_kernel(source, target)
        xx = kernels[:batch_size_s, :batch_size_s]
        yy = kernels[batch_size_s:, batch_size_s:]
        xy = kernels[:batch_size_s, batch_size_s:]
        return torch.mean(xx) + torch.mean(yy) - 2 * torch.mean(xy)

    def forward(
        self,
        source_features: torch.Tensor,
        target_features: torch.Tensor,
        source_base_soh: torch.Tensor,
        target_base_soh: torch.Tensor,
    ) -> torch.Tensor:
        if source_features.size(0) == 0 or target_features.size(0) == 0:
            return torch.tensor(0.0, device=source_features.device)
        current_soh = torch.clamp(torch.mean(target_base_soh), max=1.0)
        dynamic_weight = torch.exp(self.lambda_scale * (1.0 - current_soh))
        return dynamic_weight * self.compute_mmd(source_features, target_features)


class TemporalAttention(nn.Module):
    def __init__(self, hidden_dim: int):
        super().__init__()
        self.attention = nn.Sequential(
            nn.Linear(hidden_dim, hidden_dim // 2),
            nn.Tanh(),
            nn.Linear(hidden_dim // 2, 1),
        )

    def forward(self, lstm_outputs: torch.Tensor) -> Tuple[torch.Tensor, torch.Tensor]:
        attn_weights = F.softmax(self.attention(lstm_outputs), dim=1)
        context_vector = torch.sum(attn_weights * lstm_outputs, dim=1)
        return context_vector, attn_weights


class ResidualCRNN(nn.Module):
    def __init__(self, base_coeffs: torch.Tensor, use_physical_base: bool, use_attention: bool):
        super().__init__()
        self.base_coeffs = base_coeffs
        self.use_physical_base = use_physical_base
        self.use_attention = use_attention
        self.cnn = nn.Sequential(
            nn.Conv1d(1, 16, 3, padding=1),
            nn.BatchNorm1d(16),
            nn.ReLU(),
            nn.MaxPool1d(2),
            nn.Conv1d(16, 32, 3, padding=1),
            nn.BatchNorm1d(32),
            nn.ReLU(),
            nn.MaxPool1d(2),
        )
        self.lstm = nn.LSTM(32, 64, batch_first=True)
        self.attention = TemporalAttention(hidden_dim=64) if use_attention else None
        self.phy = nn.Sequential(nn.Linear(1, 16), nn.ReLU(), nn.Linear(16, 8), nn.ReLU())
        self.fc_soh = nn.Sequential(nn.Linear(72, 32), nn.ReLU(), nn.Dropout(0.2), nn.Linear(32, 1))

    def forward(self, x: torch.Tensor, p: torch.Tensor) -> Tuple[torch.Tensor, torch.Tensor, torch.Tensor]:
        if self.use_physical_base:
            a, b, c, d = self.base_coeffs
            base_soh = a * (p**3) + b * (p**2) + c * p + d
        else:
            base_soh = torch.zeros((x.size(0), 1), device=x.device, dtype=x.dtype)

        out = self.cnn(x.unsqueeze(1)).permute(0, 2, 1)
        lstm_out, _ = self.lstm(out)
        if self.use_attention:
            temporal_feat, _ = self.attention(lstm_out)
        else:
            temporal_feat = lstm_out[:, -1, :]
        fused_features = torch.cat((temporal_feat, self.phy(p)), dim=1)
        residual_soh = self.fc_soh(fused_features)
        pred_soh = base_soh + residual_soh if self.use_physical_base else residual_soh
        return pred_soh, fused_features, base_soh


def rmse(y_true: np.ndarray, y_pred: np.ndarray) -> float:
    return float(math.sqrt(np.mean((y_true.reshape(-1) - y_pred.reshape(-1)) ** 2)))


def mae(y_true: np.ndarray, y_pred: np.ndarray) -> float:
    return float(np.mean(np.abs(y_true.reshape(-1) - y_pred.reshape(-1))))


def mape(y_true: np.ndarray, y_pred: np.ndarray) -> float:
    denom = np.clip(np.abs(y_true.reshape(-1)), 1e-8, None)
    return float(np.mean(np.abs((y_true.reshape(-1) - y_pred.reshape(-1)) / denom)) * 100)


def macro_rmse_by_battery(y_true: np.ndarray, y_pred: np.ndarray, battery_idx: np.ndarray) -> float:
    vals = []
    for b in np.unique(battery_idx):
        mask = battery_idx == b
        if np.any(mask):
            vals.append(rmse(y_true[mask], y_pred[mask]))
    return float(np.mean(vals)) if vals else float("nan")


def evaluate(model: nn.Module, loader: DataLoader, device: torch.device) -> Tuple[Dict[str, float], pd.DataFrame]:
    model.eval()
    preds_all, trues_all, bidx_all = [], [], []
    with torch.no_grad():
        for x, y, p, b in loader:
            x = x.float().to(device)
            p = p.float().to(device)
            out, _, _ = model(x, p)
            preds_all.append(out.cpu().numpy().reshape(-1))
            trues_all.append(y.numpy().reshape(-1))
            bidx_all.append(b.numpy().reshape(-1))
    y_pred = np.concatenate(preds_all)
    y_true = np.concatenate(trues_all)
    bidx = np.concatenate(bidx_all)
    metrics = {
        "pooled_rmse": rmse(y_true, y_pred),
        "macro_rmse": macro_rmse_by_battery(y_true, y_pred, bidx),
        "mae": mae(y_true, y_pred),
        "mape_percent": mape(y_true, y_pred),
    }
    rows = []
    for b in np.unique(bidx):
        mask = bidx == b
        rows.append(
            {
                "battery_index": int(b),
                "rmse": rmse(y_true[mask], y_pred[mask]),
                "mae": mae(y_true[mask], y_pred[mask]),
                "n_points": int(np.sum(mask)),
            }
        )
    return metrics, pd.DataFrame(rows)


def split_summary(name: str, x: np.ndarray, y: np.ndarray, p: np.ndarray, bidx: np.ndarray, batch: np.ndarray) -> Dict:
    return {
        "split": name,
        "samples": int(len(x)),
        "cells": int(len(np.unique(bidx))),
        "batches": ",".join(map(str, sorted(np.unique(batch).tolist()))),
        "x_shape": "x".join(map(str, x.shape)),
        "y_min": float(np.min(y)),
        "y_max": float(np.max(y)),
        "cycle_min": float(np.min(p)),
        "cycle_max": float(np.max(p)),
    }


def train_variant(
    variant: Dict,
    data: Dict[str, Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray]],
    coeffs: np.ndarray,
    device: torch.device,
) -> Dict:
    set_seed(SEED)
    model_name = variant["model"]
    run_dir = CODEX_DIR / model_name
    run_dir.mkdir(parents=True, exist_ok=True)
    checkpoint_path = run_dir / f"{model_name}_best_by_target_adapt.pt"
    best_json = run_dir / f"{model_name}_best_metrics.json"
    result_xlsx = run_dir / f"{model_name}_results.xlsx"
    summary_xlsx = run_dir / f"{model_name}_summary.xlsx"
    history_csv = run_dir / f"{model_name}_history.csv"

    s_x, s_y, s_p, s_bidx, s_batch = data["source"]
    ta_x, ta_y, ta_p, ta_bidx, ta_batch = data["target_adapt"]
    tt_x, tt_y, tt_p, tt_bidx, tt_batch = data["target_test"]
    ds_s = HybridDataset(s_x, s_y, s_p, s_bidx)
    ds_ta = HybridDataset(ta_x, ta_y, ta_p, ta_bidx)
    ds_tt = HybridDataset(tt_x, tt_y, tt_p, tt_bidx)
    loader_s = DataLoader(ds_s, batch_size=BATCH_SIZE, shuffle=True, drop_last=True)
    loader_ta = DataLoader(ds_ta, batch_size=BATCH_SIZE, shuffle=True, drop_last=True)
    loader_ta_eval = DataLoader(ds_ta, batch_size=EVAL_BATCH_SIZE, shuffle=False)
    loader_tt_eval = DataLoader(ds_tt, batch_size=EVAL_BATCH_SIZE, shuffle=False)

    base_coeffs = torch.tensor(coeffs, dtype=torch.float32, device=device)
    model = ResidualCRNN(
        base_coeffs=base_coeffs,
        use_physical_base=variant["use_physical_base"],
        use_attention=variant["use_attention"],
    ).to(device)
    optimizer = optim.Adam(model.parameters(), lr=LEARNING_RATE)
    criterion_mse = nn.MSELoss()
    criterion_mmd = LifetimeAdaptiveMMD().to(device)

    best_target_adapt_rmse = float("inf")
    best_epoch = -1
    best_target_adapt_metrics = None
    history_rows: List[Dict] = []

    print(f"RUN={model_name}", flush=True)
    for epoch in range(1, EPOCHS + 1):
        model.train()
        if epoch == LR_DECAY_EPOCH + 1:
            for param_group in optimizer.param_groups:
                param_group["lr"] = LEARNING_RATE * LR_DECAY_FACTOR

        if epoch <= 20:
            stage = "MSE Warmup"
            mmd_weight = 0.0
        elif epoch <= 120:
            stage = "MMD Guide" if variant["use_mmd"] else "No MMD"
            mmd_weight = TARGET_MMD_WEIGHT if variant["use_mmd"] else 0.0
        else:
            stage = "Pure MSE Sprint"
            mmd_weight = 0.0

        batch_losses, mse_losses, mmd_losses = [], [], []
        for (sx, sy, sp, _), (tx, _, tp, _) in zip(loader_s, loader_ta):
            sx = sx.float().to(device)
            sy = sy.float().to(device)
            sp = sp.float().to(device)
            tx = tx.float().to(device)
            tp = tp.float().to(device)
            optimizer.zero_grad()
            pred_s, feat_s, base_s = model(sx, sp)
            _, feat_t, base_t = model(tx, tp)
            loss_mse = criterion_mse(pred_s, sy)
            loss_mmd = criterion_mmd(feat_s, feat_t, base_s, base_t) if mmd_weight > 0 else torch.tensor(0.0, device=device)
            loss = loss_mse + mmd_weight * loss_mmd
            loss.backward()
            optimizer.step()
            batch_losses.append(float(loss.item()))
            mse_losses.append(float(loss_mse.item()))
            mmd_losses.append(float(loss_mmd.item()))

        target_adapt_metrics, _ = evaluate(model, loader_ta_eval, device)
        if target_adapt_metrics["pooled_rmse"] < best_target_adapt_rmse:
            best_target_adapt_rmse = target_adapt_metrics["pooled_rmse"]
            best_epoch = epoch
            best_target_adapt_metrics = target_adapt_metrics
            torch.save(
                {
                    "epoch": epoch,
                    "model_state_dict": model.state_dict(),
                    "optimizer_state_dict": optimizer.state_dict(),
                    "base_coeffs": coeffs,
                    "selection_split": "target_adapt",
                    "target_adapt_metrics": target_adapt_metrics,
                    "variant": variant,
                    "config": {
                        "label_mode": LABEL_MODE,
                        "data_protocol": DATA_PROTOCOL,
                        "seed": SEED,
                        "epochs": EPOCHS,
                        "batch_size": BATCH_SIZE,
                        "learning_rate": LEARNING_RATE,
                        "target_mmd_weight": TARGET_MMD_WEIGHT,
                        "lambda_smooth": LAMBDA_SMOOTH,
                    },
                },
                checkpoint_path,
            )

        row = {
            "epoch": epoch,
            "stage": stage,
            "train_loss": float(np.mean(batch_losses)),
            "train_mse": float(np.mean(mse_losses)),
            "train_mmd": float(np.mean(mmd_losses)),
            "mmd_weight": float(mmd_weight),
            "lr": float(optimizer.param_groups[0]["lr"]),
            "target_adapt_pooled_rmse": target_adapt_metrics["pooled_rmse"],
            "target_adapt_macro_rmse": target_adapt_metrics["macro_rmse"],
            "target_adapt_mae": target_adapt_metrics["mae"],
            "target_adapt_mape_percent": target_adapt_metrics["mape_percent"],
            "best_epoch": best_epoch,
            "best_target_adapt_rmse": best_target_adapt_rmse,
        }
        history_rows.append(row)
        pd.DataFrame(history_rows).to_csv(history_csv, index=False, encoding="utf-8-sig")
        print(
            f"{model_name} Epoch {epoch:03d}/{EPOCHS} [{stage}] "
            f"target_adapt_rmse={row['target_adapt_pooled_rmse']:.5f} "
            f"best={best_target_adapt_rmse:.5f}@{best_epoch}",
            flush=True,
        )

    checkpoint = torch.load(checkpoint_path, map_location=device, weights_only=False)
    model.load_state_dict(checkpoint["model_state_dict"])
    target_adapt_metrics, target_adapt_cells = evaluate(model, loader_ta_eval, device)
    target_test_metrics, target_test_cells = evaluate(model, loader_tt_eval, device)

    split_rows = [
        split_summary("source", s_x, s_y, s_p, s_bidx, s_batch),
        split_summary("target_adapt", ta_x, ta_y, ta_p, ta_bidx, ta_batch),
        split_summary("target_test", tt_x, tt_y, tt_p, tt_bidx, tt_batch),
    ]
    config_row = {
        "model": model_name,
        "description": variant["description"],
        "python_executable": sys.executable,
        "device": str(device),
        "data_dir": str(DATA_DIR),
        "seed": SEED,
        "best_epoch_selected_by_target_adapt": best_epoch,
        "best_target_adapt_rmse": target_adapt_metrics["pooled_rmse"],
        "final_target_test_rmse": target_test_metrics["pooled_rmse"],
        "checkpoint_path": str(checkpoint_path),
    }
    metrics_rows = [
        {"split": "target_adapt", **target_adapt_metrics, "role": "model_selection"},
        {"split": "target_test", **target_test_metrics, "role": "final_blind_test_once"},
    ]
    compact_summary = {
        "model": model_name,
        "best_target_adapt_rmse": target_adapt_metrics["pooled_rmse"],
        "final_target_test_rmse": target_test_metrics["pooled_rmse"],
    }
    with pd.ExcelWriter(result_xlsx, engine="openpyxl") as writer:
        pd.DataFrame([config_row]).to_excel(writer, index=False, sheet_name="config")
        pd.DataFrame(split_rows).to_excel(writer, index=False, sheet_name="splits")
        pd.DataFrame(history_rows).to_excel(writer, index=False, sheet_name="history")
        pd.DataFrame(metrics_rows).to_excel(writer, index=False, sheet_name="metrics")
        target_adapt_cells.to_excel(writer, index=False, sheet_name="target_adapt_cells")
        target_test_cells.to_excel(writer, index=False, sheet_name="target_test_cells")
    pd.DataFrame([compact_summary]).to_excel(summary_xlsx, index=False, sheet_name="summary")
    best_json.write_text(
        json.dumps(
            {
                "config": config_row,
                "metrics": metrics_rows,
                "compact_summary": compact_summary,
                "target_test_cells": target_test_cells.to_dict(orient="records"),
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    del model, optimizer, criterion_mse, criterion_mmd
    if torch.cuda.is_available():
        torch.cuda.empty_cache()
    return {
        **compact_summary,
        "target_test_macro_rmse": target_test_metrics["macro_rmse"],
        "target_test_mae": target_test_metrics["mae"],
        "target_test_mape_percent": target_test_metrics["mape_percent"],
        "best_epoch": best_epoch,
        "description": variant["description"],
        "run_dir": str(run_dir),
    }


def load_existing_final_retained() -> List[Dict]:
    rows = []
    summary_path = CODEX_DIR / "final-retained-v1" / "final-retained-v1_summary.xlsx"
    if summary_path.exists():
        df = pd.read_excel(summary_path)
        rows.extend(df.to_dict(orient="records"))
    return rows


def update_global_summary(rows: List[Dict]) -> pd.DataFrame:
    final_rows = []
    for row in rows:
        final_rows.append(
            {
                "model": row["model"],
                "best_target_adapt_rmse": row["best_target_adapt_rmse"],
                "final_target_test_rmse": row["final_target_test_rmse"],
            }
        )
    df = pd.DataFrame(final_rows, columns=["model", "best_target_adapt_rmse", "final_target_test_rmse"])
    df.to_excel(SUMMARY_XLSX, index=False)
    return df


def set_run_font(run, size=11, bold=False, font="宋体"):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True, font="黑体")


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(22)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=11)


def add_table(doc: Document, headers: List[str], rows: List[List]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        run = table.rows[0].cells[i].paragraphs[0].add_run(str(h))
        set_run_font(run, size=9.5, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            run = cells[i].paragraphs[0].add_run(str(value))
            set_run_font(run, size=9.5)
    doc.add_paragraph()


def build_report(result_rows: List[Dict], summary_df: pd.DataFrame) -> None:
    SUITE_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = 1
    run = title.add_run("严格方案 A 消融实验汇报")
    set_run_font(run, size=16, bold=True, font="黑体")
    add_para(
        doc,
        "本次实验在 processed_clean_fixed_split 协议下执行，Batch 1–8 作为 source，Batch 9 作为 target_adapt 并用于模型选择，Batch 10 仅作为最终 blind-test 评估一次。所有消融均基于最终保留模型设置，不重新引入当前形式的一阶差分 smoothness penalty。",
    )
    add_heading(doc, "1. 总体结果", 1)
    add_table(
        doc,
        ["模型", "target_adapt 最佳 RMSE", "target_test 最终 RMSE"],
        [
            [row["model"], f"{row['best_target_adapt_rmse']:.6f}", f"{row['final_target_test_rmse']:.6f}"]
            for row in summary_df.to_dict(orient="records")
        ],
    )
    add_heading(doc, "2. 消融结果解读", 1)
    by_model = {row["model"]: row for row in result_rows}
    final_df = summary_df[summary_df["model"] == "final-retained-v1"]
    final_rmse = float(final_df.iloc[0]["final_target_test_rmse"]) if len(final_df) else None
    for model in ["no-physical-base-v1", "no-mmd-v1", "no-attention-v1"]:
        row = by_model.get(model)
        if not row:
            continue
        delta = row["final_target_test_rmse"] - final_rmse if final_rmse is not None else float("nan")
        add_para(
            doc,
            f"{model}：{row['description']}该模型在 target_test 上的 pooled RMSE 为 {row['final_target_test_rmse']:.6f}，"
            f"相对于 final-retained-v1 的变化为 {delta:+.6f}。该结果用于判断对应模块在严格 blind-test 协议下的实际贡献。",
        )
    add_heading(doc, "3. 当前结论与注意事项", 1)
    add_para(
        doc,
        "若 no-Physical-Base 的 RMSE 明显高于 final-retained-v1，则可继续支持“物理退化基座是核心模块”的论点；若 no-MMD 或 no-Attention 的差异较小，则论文中应使用更克制的表述，将其写作跨域增强或辅助表征模块，而不是绝对化地描述为普遍提升模块。",
    )
    add_para(
        doc,
        "本次实验已经避免使用 target_test 进行 checkpoint selection，因此比早期 test-guided 结果更适合作为论文最终主表的候选结果。后续如果时间允许，建议继续补充 seed=7 和 seed=123 的重复实验，并报告均值±标准差。",
    )
    doc.save(REPORT_DOCX)


def main() -> None:
    CODEX_DIR.mkdir(exist_ok=True)
    SUITE_DIR.mkdir(exist_ok=True)
    set_seed(SEED)
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    required = [
        DATA_DIR / "clean_source_x.npy",
        DATA_DIR / "clean_source_y_nominal.npy",
        DATA_DIR / "clean_target_adapt_x.npy",
        DATA_DIR / "clean_target_adapt_y_nominal.npy",
        DATA_DIR / "clean_target_test_x.npy",
        DATA_DIR / "clean_target_test_y_nominal.npy",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("processed_clean is incomplete: " + "; ".join(missing))

    print(f"python={sys.executable}", flush=True)
    print(f"device={device}", flush=True)
    print(f"data_dir={DATA_DIR}", flush=True)
    data = {
        "source": load_split("source", LABEL_MODE),
        "target_adapt": load_split("target_adapt", LABEL_MODE),
        "target_test": load_split("target_test", LABEL_MODE),
    }
    s_x, s_y, s_p, _, _ = data["source"]
    coeffs = np.polyfit((s_p.flatten() / GLOBAL_MAX_CYCLE), s_y.flatten(), 3)
    print(f"base_coeffs={coeffs}", flush=True)

    result_rows = []
    existing_rows = load_existing_final_retained()
    result_rows.extend(existing_rows)
    for variant in VARIANTS:
        result = train_variant(variant, data, coeffs, device)
        result_rows.append(result)

    summary_df = update_global_summary(result_rows)
    build_report(result_rows, summary_df)
    print("ABLATION_SUITE_FINAL_SUMMARY", summary_df.to_json(orient="records", force_ascii=False), flush=True)
    print(f"REPORT={REPORT_DOCX}", flush=True)


if __name__ == "__main__":
    main()
