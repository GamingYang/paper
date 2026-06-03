import json
import math
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"C:\Users\ROG\Desktop\paper")
CODEX_DIR = ROOT / "codex"
OUT_DIR = CODEX_DIR / "paper-draft-v3-20260425"
FIG_DIR = OUT_DIR / "figures"

DOCX_OUT = OUT_DIR / "method_experiment_draft_v3.docx"
XLSX_OUT = OUT_DIR / "manuscript_v3_metrics_summary.xlsx"
JSON_OUT = OUT_DIR / "manuscript_v3_generation_summary_20260425.json"

V2_DIR = CODEX_DIR / "paper-draft-v2-20260420"
V2_FIG_DIR = V2_DIR / "figures"

FINAL_RESULTS_XLSX = CODEX_DIR / "final-retained-v1" / "final-retained-v1_results.xlsx"
FINAL_PER_BATTERY_XLSX = CODEX_DIR / "paper-update-20260420" / "per_battery_ablation_comparison_20260420.xlsx"
NO_PHYSICAL_JSON = CODEX_DIR / "no-physical-base-v1" / "no-physical-base-v1_best_metrics.json"
NO_MMD_JSON = CODEX_DIR / "no-mmd-v1" / "no-mmd-v1_best_metrics.json"
NO_ATT_JSON = CODEX_DIR / "no-attention-v1" / "no-attention-v1_best_metrics.json"
EXTERNAL_V1_XLSX = CODEX_DIR / "external-baseline-suite-v1" / "external_baseline_suite_v1_summary.xlsx"
KERNEL_V2_XLSX = CODEX_DIR / "kernel-baseline-suite-v2" / "kernel_baseline_suite_v2_summary.xlsx"
OLD_INTERNAL_SUMMARY_XLSX = CODEX_DIR / "model_runs_summary.xlsx"


def ensure_dirs() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)


def configure_doc(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_paragraph(doc: Document, text: str, *, bold: bool = False, italic: bool = False, center: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_table_from_df(doc: Document, df: pd.DataFrame, title: str) -> None:
    add_paragraph(doc, title, bold=True)
    view = df.copy()
    for col in view.columns:
        if pd.api.types.is_float_dtype(view[col]):
            view[col] = view[col].map(lambda x: f"{float(x):.6f}" if pd.notna(x) else "")
    table = doc.add_table(rows=1, cols=len(view.columns))
    table.style = "Table Grid"
    for idx, col in enumerate(view.columns):
        table.rows[0].cells[idx].text = str(col)
    for _, row in view.iterrows():
        cells = table.add_row().cells
        for idx, val in enumerate(row.tolist()):
            cells[idx].text = str(val)


def add_figure(doc: Document, image_path: Path, caption: str, width_inches: float = 6.2) -> None:
    doc.add_picture(str(image_path), width=Inches(width_inches))
    add_paragraph(doc, caption, italic=True, center=True)


def load_internal_final() -> dict:
    metrics_df = pd.read_excel(FINAL_RESULTS_XLSX, sheet_name="metrics")
    cfg_df = pd.read_excel(FINAL_RESULTS_XLSX, sheet_name="config")
    target_dev = metrics_df.loc[metrics_df["split"] == "target_adapt"].iloc[0]
    target_test = metrics_df.loc[metrics_df["split"] == "target_test"].iloc[0]
    return {
        "category": "Retained model",
        "group": "existing_internal",
        "model": "final-retained-v1",
        "label": "Final retained model",
        "target_dev_pooled_rmse": float(target_dev["pooled_rmse"]),
        "target_test_pooled_rmse": float(target_test["pooled_rmse"]),
        "target_test_macro_rmse": float(target_test["macro_rmse"]),
        "target_test_mae": float(target_test["mae"]),
        "target_test_mape_percent": float(target_test["mape_percent"]),
        "notes": "Physics base + residual temporal network + Attention + adaptive MMD",
        "source_file": str(FINAL_RESULTS_XLSX),
        "best_epoch": int(cfg_df.loc[0, "best_epoch_selected_by_target_adapt"]),
    }


def load_internal_from_json(path: Path, label: str, notes: str) -> dict:
    payload = json.loads(path.read_text(encoding="utf-8"))
    target_dev = next(item for item in payload["metrics"] if item["split"] == "target_adapt")
    target_test = next(item for item in payload["metrics"] if item["split"] == "target_test")
    return {
        "category": "Internal ablation",
        "group": "existing_internal",
        "model": payload["config"]["model"],
        "label": label,
        "target_dev_pooled_rmse": float(target_dev["pooled_rmse"]),
        "target_test_pooled_rmse": float(target_test["pooled_rmse"]),
        "target_test_macro_rmse": float(target_test["macro_rmse"]),
        "target_test_mae": float(target_test["mae"]),
        "target_test_mape_percent": float(target_test["mape_percent"]),
        "notes": notes,
        "source_file": str(path),
        "best_epoch": int(payload["config"]["best_epoch_selected_by_target_adapt"]),
    }


def load_external_v1() -> pd.DataFrame:
    df = pd.read_excel(EXTERNAL_V1_XLSX, sheet_name="new_external").copy()
    category_map = {
        "physical-only-v1": "External baseline",
        "lstm-baseline-v1": "External baseline",
        "gru-baseline-v1": "External baseline",
        "cnn-lstm-baseline-v1": "External baseline",
    }
    note_map = {
        "physical-only-v1": "Pure cubic physical degradation base",
        "lstm-baseline-v1": "Pure data-driven recurrent baseline",
        "gru-baseline-v1": "Pure data-driven gated recurrent baseline",
        "cnn-lstm-baseline-v1": "Convolutional-recurrent baseline without explicit physics or MMD",
    }
    df["category"] = df["model"].map(category_map)
    df["notes"] = df["model"].map(note_map)
    df["source_file"] = str(EXTERNAL_V1_XLSX)
    df = df.rename(
        columns={
            "target_dev_rmse": "target_dev_pooled_rmse",
            "target_test_rmse": "target_test_pooled_rmse",
        }
    )
    keep_cols = [
        "category",
        "group",
        "model",
        "label",
        "target_dev_pooled_rmse",
        "target_test_pooled_rmse",
        "target_test_macro_rmse",
        "target_test_mae",
        "target_test_mape_percent",
        "notes",
        "source_file",
    ]
    return df[keep_cols].copy()


def load_kernel_v2() -> pd.DataFrame:
    df = pd.read_excel(KERNEL_V2_XLSX, sheet_name="new_kernel").copy()
    note_map = {
        "pca-gpr-v1": "Kernel statistical baseline with source-side stratified subsampling and PCA compression",
        "pca-rbf-svr-v1": "RBF-SVR baseline with source-side stratified subsampling and PCA compression",
        "pca-linearsvr-v1": "LinearSVR baseline with source-side stratified subsampling and PCA compression",
    }
    df["category"] = "Kernel baseline"
    df["notes"] = df["model"].map(note_map)
    df["source_file"] = str(KERNEL_V2_XLSX)
    df = df.rename(
        columns={
            "target_dev_rmse": "target_dev_pooled_rmse",
            "target_test_rmse": "target_test_pooled_rmse",
        }
    )
    keep_cols = [
        "category",
        "group",
        "model",
        "label",
        "target_dev_pooled_rmse",
        "target_test_pooled_rmse",
        "target_test_macro_rmse",
        "target_test_mae",
        "target_test_mape_percent",
        "notes",
        "source_file",
        "target_dev_macro_rmse",
        "target_dev_mae",
        "source_subset_size",
        "pca_components",
        "best_params_json",
    ]
    return df[keep_cols].copy()


def make_combined_summary() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    internal_rows = [
        load_internal_final(),
        load_internal_from_json(
            NO_PHYSICAL_JSON,
            "w/o physical base",
            "Remove cubic physical degradation base and directly regress SOH from residual network",
        ),
        load_internal_from_json(
            NO_MMD_JSON,
            "w/o MMD",
            "Disable adaptive MMD while retaining physics base and Attention",
        ),
        load_internal_from_json(
            NO_ATT_JSON,
            "w/o Attention",
            "Remove Attention and use last LSTM hidden state for temporal aggregation",
        ),
    ]
    internal_df = pd.DataFrame(internal_rows)
    external_df = load_external_v1()
    kernel_df = load_kernel_v2()
    combined_df = pd.concat([internal_df, external_df, kernel_df], ignore_index=True, sort=False)
    combined_df = combined_df.sort_values("target_test_pooled_rmse", ascending=True).reset_index(drop=True)
    return combined_df, internal_df, external_df, kernel_df


def copy_figures() -> list[Path]:
    figure_names = [
        "final-retained-v1_batch10_2x4_clean_abs_error_bars_v2.png",
        "per_battery_rmse_comparison_bar_v2.png",
        "no-physical-base-v1_batch10_2x4_clean_abs_error_bars_v2.png",
        "no-mmd-v1_batch10_2x4_clean_abs_error_bars_v2.png",
        "no-attention-v1_batch10_2x4_clean_abs_error_bars_v2.png",
    ]
    copied = []
    for name in figure_names:
        src = V2_FIG_DIR / name
        dst = FIG_DIR / name
        shutil.copy2(src, dst)
        copied.append(dst)
    return copied


def relative_change(base: float, variant: float) -> float:
    return (variant - base) / base * 100.0


def reduction_vs_baseline(baseline: float, improved: float) -> float:
    return (baseline - improved) / baseline * 100.0


def validate(combined_df: pd.DataFrame, internal_df: pd.DataFrame) -> dict:
    internal_no_nan = bool(internal_df["target_test_macro_rmse"].notna().all())
    reference_checks = {
        "final_target_test_rmse": math.isclose(
            float(combined_df.loc[combined_df["model"] == "final-retained-v1", "target_test_pooled_rmse"].iloc[0]),
            0.009499246123112539,
            rel_tol=0.0,
            abs_tol=1e-9,
        ),
        "final_macro_rmse": math.isclose(
            float(combined_df.loc[combined_df["model"] == "final-retained-v1", "target_test_macro_rmse"].iloc[0]),
            0.008301,
            rel_tol=0.0,
            abs_tol=1e-6,
        ),
        "no_physical_macro_rmse": math.isclose(
            float(combined_df.loc[combined_df["model"] == "no-physical-base-v1", "target_test_macro_rmse"].iloc[0]),
            0.019506288473427214,
            rel_tol=0.0,
            abs_tol=1e-9,
        ),
        "no_mmd_macro_rmse": math.isclose(
            float(combined_df.loc[combined_df["model"] == "no-mmd-v1", "target_test_macro_rmse"].iloc[0]),
            0.008721296235627948,
            rel_tol=0.0,
            abs_tol=1e-9,
        ),
        "no_attention_macro_rmse": math.isclose(
            float(combined_df.loc[combined_df["model"] == "no-attention-v1", "target_test_macro_rmse"].iloc[0]),
            0.010534793547796156,
            rel_tol=0.0,
            abs_tol=1e-9,
        ),
        "cnn_lstm_rmse": math.isclose(
            float(combined_df.loc[combined_df["model"] == "cnn-lstm-baseline-v1", "target_test_pooled_rmse"].iloc[0]),
            0.016075,
            rel_tol=0.0,
            abs_tol=1e-6,
        ),
        "pca_gpr_rmse": math.isclose(
            float(combined_df.loc[combined_df["model"] == "pca-gpr-v1", "target_test_pooled_rmse"].iloc[0]),
            0.023108,
            rel_tol=0.0,
            abs_tol=1e-6,
        ),
    }
    ranking = combined_df[["model", "target_test_pooled_rmse"]].copy()
    return {
        "internal_macro_rmse_complete": internal_no_nan,
        "reference_checks": reference_checks,
        "ranking_preview": ranking.to_dict(orient="records"),
    }


def write_excel(
    combined_df: pd.DataFrame,
    internal_df: pd.DataFrame,
    external_df: pd.DataFrame,
    kernel_df: pd.DataFrame,
    validation_payload: dict,
) -> None:
    display_df = combined_df[
        [
            "category",
            "label",
            "target_dev_pooled_rmse",
            "target_test_pooled_rmse",
            "target_test_macro_rmse",
            "target_test_mae",
            "target_test_mape_percent",
        ]
    ].copy()
    display_df = display_df.rename(
        columns={
            "category": "Category",
            "label": "Model",
            "target_dev_pooled_rmse": "Target-dev pooled RMSE",
            "target_test_pooled_rmse": "Batch10 pooled RMSE",
            "target_test_macro_rmse": "Batch10 macro RMSE",
            "target_test_mae": "Batch10 MAE",
            "target_test_mape_percent": "Batch10 MAPE (%)",
        }
    )
    validation_df = pd.DataFrame(
        {
            "check": list(validation_payload["reference_checks"].keys()) + ["internal_macro_rmse_complete"],
            "passed": list(validation_payload["reference_checks"].values()) + [validation_payload["internal_macro_rmse_complete"]],
        }
    )
    with pd.ExcelWriter(XLSX_OUT, engine="openpyxl") as writer:
        display_df.to_excel(writer, sheet_name="combined_summary", index=False)
        internal_df.to_excel(writer, sheet_name="internal_backfilled", index=False)
        external_df.to_excel(writer, sheet_name="external_baselines", index=False)
        kernel_df.to_excel(writer, sheet_name="kernel_baselines", index=False)
        validation_df.to_excel(writer, sheet_name="validation_check", index=False)


def build_doc(
    combined_df: pd.DataFrame,
    internal_df: pd.DataFrame,
    external_df: pd.DataFrame,
    kernel_df: pd.DataFrame,
    copied_figures: list[Path],
) -> None:
    doc = Document()
    configure_doc(doc)

    final_row = combined_df.loc[combined_df["model"] == "final-retained-v1"].iloc[0]
    no_physical_row = combined_df.loc[combined_df["model"] == "no-physical-base-v1"].iloc[0]
    no_mmd_row = combined_df.loc[combined_df["model"] == "no-mmd-v1"].iloc[0]
    no_attention_row = combined_df.loc[combined_df["model"] == "no-attention-v1"].iloc[0]
    cnn_lstm_row = combined_df.loc[combined_df["model"] == "cnn-lstm-baseline-v1"].iloc[0]
    pca_gpr_row = combined_df.loc[combined_df["model"] == "pca-gpr-v1"].iloc[0]
    pca_rbf_row = combined_df.loc[combined_df["model"] == "pca-rbf-svr-v1"].iloc[0]
    pca_linear_row = combined_df.loc[combined_df["model"] == "pca-linearsvr-v1"].iloc[0]

    no_physical_change = relative_change(final_row["target_test_pooled_rmse"], no_physical_row["target_test_pooled_rmse"])
    no_mmd_change = relative_change(final_row["target_test_pooled_rmse"], no_mmd_row["target_test_pooled_rmse"])
    no_attention_change = relative_change(final_row["target_test_pooled_rmse"], no_attention_row["target_test_pooled_rmse"])
    vs_cnn_lstm = reduction_vs_baseline(cnn_lstm_row["target_test_pooled_rmse"], final_row["target_test_pooled_rmse"])
    vs_pca_gpr = reduction_vs_baseline(pca_gpr_row["target_test_pooled_rmse"], final_row["target_test_pooled_rmse"])

    doc.add_heading("方法、实验设置与结果分析初稿 v3", level=0)

    doc.add_heading("1 Methodology", level=1)
    doc.add_heading("1.1 Overall Framework", level=2)
    add_paragraph(
        doc,
        "本文构建一种面向跨批次锂离子电池 SOH 估计的物理引导混合迁移学习框架。整体方法由全局物理退化基座、残差时序建模网络和生命周期自适应域对齐约束组成。"
        "三者承担不同层级的功能：物理基座刻画跨电池共享的长期退化趋势，残差时序网络学习围绕该基准线的个体化局部偏差，MMD 对齐项在训练阶段缓解 source domain 与 target-domain development set 之间的特征分布差异。"
    )
    add_paragraph(doc, "ŷ = f_phy(τ) + f_res(x, τ)", italic=True, center=True)
    add_paragraph(
        doc,
        f"该分解使神经网络不必独立承担完整 SOH 曲线拟合，而是集中学习难以由低阶退化函数表达的局部动态波动。"
        f"严格 blind-test 结果显示，移除物理基座后 Batch 10 pooled RMSE 从 {final_row['target_test_pooled_rmse']:.6f} 上升至 {no_physical_row['target_test_pooled_rmse']:.6f}，"
        f"相对恶化 {no_physical_change:.2f}%，说明宏观退化基准线是当前方法中最关键的结构支撑。"
    )

    doc.add_heading("1.2 Global Physical Degradation Base", level=2)
    add_paragraph(
        doc,
        "电池容量衰减具有明显的长期趋势性。为避免纯数据驱动网络将宏观退化、批次差异和局部噪声混合到同一隐空间中，本文首先使用低复杂度全局物理退化基座为所有样本建立共享参考曲线。"
        "当前实现采用三次多项式形式："
    )
    add_paragraph(doc, "f_phy(τ) = aτ^3 + bτ^2 + cτ + d", italic=True, center=True)
    add_paragraph(
        doc,
        "其中 τ 为归一化循环位置，a、b、c、d 为由 source domain 数据估计得到的全局系数。该三次多项式并不被视为复杂原创结构，而是在当前任务中承担明确功能的物理退化基准线。"
        "它提供宏观趋势先验，使残差网络主要处理个体化偏差。"
    )

    doc.add_heading("1.3 Residual Temporal Modeling", level=2)
    add_paragraph(
        doc,
        "在物理基座给出长期退化趋势后，仍需对不同电池在局部循环阶段表现出的个体化波动进行建模。本文采用由 CNN、LSTM 和 Attention 组成的残差时序网络。"
        "CNN 用于提取局部序列模式，LSTM 建模窗口内时序依赖，Attention 对不同时刻隐状态进行加权聚合，以形成更集中的残差表征。"
    )
    add_paragraph(
        doc,
        f"Attention 在本文中定位为辅助表征模块，而不是核心原创机制。严格协议下，移除 Attention 后 Batch 10 pooled RMSE 由 {final_row['target_test_pooled_rmse']:.6f} 上升至 {no_attention_row['target_test_pooled_rmse']:.6f}，"
        f"相对恶化 {no_attention_change:.2f}%；对应 macro RMSE 由 {final_row['target_test_macro_rmse']:.6f} 上升至 {no_attention_row['target_test_macro_rmse']:.6f}。"
        "因此论文中应将其描述为显著有效的辅助时序表征模块，而不宜夸大为主要创新来源。"
    )

    doc.add_heading("1.4 Adaptive Domain Alignment", level=2)
    add_paragraph(
        doc,
        "跨批次 SOH 估计面临 source batches 与目标批次之间的分布偏移。为提高目标域泛化能力，本文在训练阶段引入生命周期自适应 MMD 对齐项，使 source domain 与 target-domain development set 的中间特征分布保持适度接近。训练目标可概括为："
    )
    add_paragraph(doc, "L = L_reg + λ(τ)L_MMD", italic=True, center=True)
    add_paragraph(
        doc,
        "其中 source 与 target-domain development samples 的中间表征用于训练阶段域对齐，λ 表示随目标域退化阶段变化的自适应权重。"
        "需要强调的是，MMD 只在训练阶段使用 Batch 9 target-domain development set 进行对齐；Batch 10 始终保留为最终 blind-test，不参与训练、调参或 checkpoint selection。"
    )
    add_paragraph(
        doc,
        f"当前消融结果显示，移除 MMD 后 Batch 10 pooled RMSE 为 {no_mmd_row['target_test_pooled_rmse']:.6f}，相较完整模型上升 {no_mmd_change:.2f}%；"
        f"macro RMSE 由 {final_row['target_test_macro_rmse']:.6f} 轻微上升至 {no_mmd_row['target_test_macro_rmse']:.6f}。"
        "因此 MMD 更适合被表述为轻量但正向的跨域增强模块。"
    )

    doc.add_heading("1.5 Optimization and Model Selection", level=2)
    add_paragraph(
        doc,
        "训练过程遵循严格 Scheme A 协议。模型首先在 source domain 上进行监督回归学习，随后结合 target-domain development set 的特征对齐约束，并根据该 development set 上的 RMSE 选择最佳 checkpoint。"
        "最终性能仅在 Batch 10 blind-test set 上评估一次。该协议避免了使用测试集选择 checkpoint 的信息泄露风险，更适合作为论文主结果的实验依据。当前 smoothness penalty 不纳入最终保留模型。"
    )

    doc.add_heading("2 Experimental Settings", level=1)
    doc.add_heading("2.1 Dataset and Splitting Protocol", level=2)
    add_paragraph(
        doc,
        "实验采用商业 LFP/graphite 锂离子电池寿命数据集，当前预处理版本包含 77 个电池单体。固定划分协议如下：Batch 1-8 作为 source domain，用于监督训练；"
        "Batch 9 作为 target-domain development set，用于目标域自适应和 checkpoint selection；Batch 10 作为 target-domain blind-test set，仅用于最终泛化评估。"
    )
    add_paragraph(
        doc,
        "当前数据规模为 source domain 115256 个样本、61 个电池；target-domain development set 16310 个样本、8 个电池；target-domain blind-test set 14556 个样本、8 个电池。"
        "输入窗口长度为 100，标签采用 nominal SOH，循环位置归一化后输入物理基座和位置表征分支。"
    )

    doc.add_heading("2.2 Compared Models", level=2)
    add_paragraph(doc, "Final retained model：包含三次物理退化基座、残差时序网络、Attention 和生命周期自适应 MMD。")
    add_paragraph(doc, "w/o physical base：移除三次物理退化基座，由残差网络直接预测 SOH。")
    add_paragraph(doc, "w/o MMD：关闭训练阶段 MMD 对齐，保留物理基座和 Attention。")
    add_paragraph(doc, "w/o Attention：移除 Attention，使用 LSTM 最后时刻隐状态作为时序表征。")
    add_paragraph(doc, "Pure cubic physical base：仅保留全局三次退化基座，用于刻画统一的宏观衰减趋势。")
    add_paragraph(doc, "LSTM / GRU / CNN-LSTM baselines：作为常见纯数据驱动时序方法，对比不引入显式物理基座和域对齐时的泛化水平。")
    add_paragraph(
        doc,
        "PCA-GPR / PCA-RBF-SVR / PCA-LinearSVR baselines：作为补充的核方法与统计学习基线。考虑到 source domain 具有 10 万级窗口样本，"
        "本轮核方法仅在 source 端执行分层下采样与 PCA 压缩，再使用 Batch 9 development set 完成候选选择；Batch 10 不参与任何超参数选择。"
    )

    doc.add_heading("2.3 Training and Evaluation Metrics", level=2)
    add_paragraph(
        doc,
        "可训练神经网络模型在相同数据划分下训练 150 个 epoch，batch size 为 128，初始学习率为 1e-3，并在后期进行学习率衰减。最佳模型由 target-domain development set RMSE 选择。"
        "对于不依赖 epoch 训练的核方法与统计学习基线，则在相同 target-domain development set 上完成候选参数选择。"
    )
    add_paragraph(
        doc,
        "主表同时报告 target-dev pooled RMSE、Batch 10 pooled RMSE、Batch 10 macro RMSE、MAE 和 MAPE。"
        "其中 pooled RMSE 反映所有测试样本的整体误差，macro RMSE 先计算每个电池单体 RMSE 后再平均，可降低样本数差异造成的偏置。"
        "为避免信息缺失，本文已将内部消融的 macro RMSE、MAE 和 MAPE 从各自原始结果文件中统一回填至总表。"
    )

    doc.add_heading("2.4 Visualization for Blind-test Prediction", level=2)
    add_paragraph(
        doc,
        "为展示模型在 Batch 10 不同电池单体上的预测行为，本文绘制 2 × 4 SOH 轨迹图。每个子图包含 measured SOH、显示层面平滑后的 predicted SOH、"
        "基于 Batch 9 residual standard deviation 构造的经验区间带，以及非负绝对误差矩形条。需要说明的是，平滑仅用于图像展示，"
        "所有 RMSE、MAE 和 MAPE 均基于原始预测值计算。"
    )

    add_figure(doc, copied_figures[0], "Figure 1. Batch 10 blind-test SOH prediction of the final retained model.")
    add_figure(doc, copied_figures[1], "Figure 2. Per-battery RMSE comparison among the retained model and ablation variants.")
    add_figure(doc, copied_figures[2], "Figure 3. Batch 10 blind-test SOH prediction of w/o physical base.")
    add_figure(doc, copied_figures[3], "Figure 4. Batch 10 blind-test SOH prediction of w/o MMD.")
    add_figure(doc, copied_figures[4], "Figure 5. Batch 10 blind-test SOH prediction of w/o Attention.")

    doc.add_heading("3 Results and Comparative Analysis", level=1)
    doc.add_heading("3.1 Overall Comparison under Strict Blind-test Protocol", level=2)
    add_paragraph(
        doc,
        f"表 1 汇总了当前已完成的 retained model、内部消融、常见外部基线以及新增核方法基线结果。"
        f"在全部已完成模型中，最终保留模型取得最优 blind-test 性能，其 Batch 10 pooled RMSE 为 {final_row['target_test_pooled_rmse']:.6f}，"
        f"macro RMSE 为 {final_row['target_test_macro_rmse']:.6f}。在既有外部基线中，CNN-LSTM 的 pooled RMSE 为 {cnn_lstm_row['target_test_pooled_rmse']:.6f}，"
        f"为当前最强外部基线；在今晚新增的核方法基线中，PCA-GPR 的 pooled RMSE 为 {pca_gpr_row['target_test_pooled_rmse']:.6f}，为核方法组中表现最优者。"
    )
    add_paragraph(
        doc,
        f"与 CNN-LSTM 相比，最终模型的 pooled RMSE 进一步下降 {vs_cnn_lstm:.2f}%；与本轮最优核方法 PCA-GPR 相比，"
        f"最终模型的 pooled RMSE 进一步下降 {vs_pca_gpr:.2f}%。这说明在严格 cross-batch blind-test 条件下，"
        "单纯依赖卷积-循环堆叠或浅层核映射仍不足以充分刻画寿命退化中的宏观趋势、局部波动与跨批次分布偏移。"
    )
    result_table = combined_df[
        [
            "category",
            "label",
            "target_dev_pooled_rmse",
            "target_test_pooled_rmse",
            "target_test_macro_rmse",
            "target_test_mae",
            "target_test_mape_percent",
        ]
    ].copy()
    result_table = result_table.rename(
        columns={
            "category": "Category",
            "label": "Model",
            "target_dev_pooled_rmse": "Target-dev pooled RMSE",
            "target_test_pooled_rmse": "Batch10 pooled RMSE",
            "target_test_macro_rmse": "Batch10 macro RMSE",
            "target_test_mae": "Batch10 MAE",
            "target_test_mape_percent": "Batch10 MAPE (%)",
        }
    )
    add_table_from_df(doc, result_table, "Table 1. Overall comparison under the strict blind-test protocol.")

    doc.add_heading("3.2 Internal Ablation Analysis", level=2)
    add_paragraph(
        doc,
        f"内部消融结果表明，物理退化基座是当前方法中最关键的组成。移除物理基座后，Batch 10 pooled RMSE 从 {final_row['target_test_pooled_rmse']:.6f} "
        f"上升至 {no_physical_row['target_test_pooled_rmse']:.6f}，macro RMSE 从 {final_row['target_test_macro_rmse']:.6f} 上升至 {no_physical_row['target_test_macro_rmse']:.6f}。"
        "该结果说明仅依赖残差网络直接回归 SOH 时，模型需要同时承担宏观退化趋势与个体化波动的双重建模任务，跨批次泛化稳定性明显下降。"
    )
    add_paragraph(
        doc,
        f"移除 Attention 后，Batch 10 pooled RMSE 为 {no_attention_row['target_test_pooled_rmse']:.6f}，相较完整模型有可见退化；对应 macro RMSE 为 {no_attention_row['target_test_macro_rmse']:.6f}。"
        "这一结果支持将 Attention 定位为有效的辅助时序表征模块，但其作用仍从属于整体残差建模框架，而非方法的唯一创新来源。"
    )
    add_paragraph(
        doc,
        f"移除 MMD 后，Batch 10 pooled RMSE 为 {no_mmd_row['target_test_pooled_rmse']:.6f}，与完整模型的差距相对较小，但整体仍表现为负向变化。"
        f"同时，macro RMSE 从 {final_row['target_test_macro_rmse']:.6f} 轻微上升至 {no_mmd_row['target_test_macro_rmse']:.6f}。"
        "因此，MMD 更适合作为轻量但正向的跨域增强模块来表述，其作用在于降低 source 与 target-domain development set 之间的表征偏移，而不是单独主导性能提升。"
    )

    doc.add_heading("3.3 External Baselines and Kernel Statistical Models", level=2)
    add_paragraph(
        doc,
        f"在外部对比基线中，Pure cubic physical base、LSTM 和 GRU 的 Batch 10 pooled RMSE 分别为 0.032649、0.032742 和 0.033109，"
        f"均明显高于最终模型。CNN-LSTM 取得 {cnn_lstm_row['target_test_pooled_rmse']:.6f} 的 pooled RMSE 和 {cnn_lstm_row['target_test_macro_rmse']:.6f} 的 macro RMSE，"
        "说明在常见纯数据驱动结构中，卷积与循环结构的组合能够提升局部特征提取与序列建模能力，但仍难以在严格跨批次 blind-test 中达到最终保留模型的表现。"
    )
    add_paragraph(
        doc,
        f"今晚新增的核方法基线中，PCA-GPR、PCA-RBF-SVR 和 PCA-LinearSVR 的 Batch 10 pooled RMSE 分别为 {pca_gpr_row['target_test_pooled_rmse']:.6f}、"
        f"{pca_rbf_row['target_test_pooled_rmse']:.6f} 和 {pca_linear_row['target_test_pooled_rmse']:.6f}。"
        f"其中 PCA-GPR 的性能优于纯三次物理基线以及 LSTM/GRU 等弱外部基线，但仍显著落后于最终模型和 CNN-LSTM。"
        "这一现象表明，在当前 10 万级 source 窗口样本与跨批次分布偏移并存的任务中，浅层核映射可以学习局部非线性关系，但仍不足以稳定覆盖宏观退化先验与个体化残差修正。"
    )
    add_paragraph(
        doc,
        "需要说明的是，为保证 GPR 与 SVR 类方法在当前样本规模下具有可计算性，本文对核方法仅在 source 端执行分层下采样与 PCA 压缩。"
        "该处理不改变 Batch 9 仅用于开发集选择、Batch 10 仅用于最终盲测的评估协议，因此不会引入 target_test 信息泄露。"
    )

    doc.add_heading("3.4 Discussion of Model Contribution", level=2)
    add_paragraph(
        doc,
        "综合内部消融与外部对比结果可以看到，当前 retained model 的优势并不来自单一模块叠加，而是来自不同组成部分之间的明确分工。"
        "物理基座负责提供统一的宏观退化参考，残差时序网络负责学习围绕基准线的个体化波动，Attention 负责辅助聚合时序信息，MMD 则在训练阶段缓解跨批次特征偏移。"
    )
    add_paragraph(
        doc,
        "如果仅保留物理基座，模型难以覆盖不同电池在局部阶段上的非线性波动；如果完全移除物理基座，则残差网络需要同时承担长期趋势与局部扰动的建模任务，"
        "导致 blind-test 性能明显恶化；如果只依赖传统循环网络或浅层核方法，则模型能够学习部分序列模式或非线性映射，但对跨批次泛化的支撑仍不充分。"
    )
    add_paragraph(
        doc,
        "因此，现有证据链支持将本文方法概括为一种以物理退化基座为主线、以残差时序建模为核心数据驱动补偿、并通过训练阶段轻量域对齐增强目标域泛化的混合框架。"
        "从审稿视角看，这种表述较能与当前实验事实保持一致，也更有利于支撑方法的科学性与创新性。"
    )

    doc.save(DOCX_OUT)


def write_json_summary(
    combined_df: pd.DataFrame,
    validation_payload: dict,
    copied_figures: list[Path],
) -> None:
    payload = {
        "outputs": {
            "docx": str(DOCX_OUT),
            "xlsx": str(XLSX_OUT),
            "figures_dir": str(FIG_DIR),
        },
        "source_files": {
            "old_internal_summary": str(OLD_INTERNAL_SUMMARY_XLSX),
            "final_results": str(FINAL_RESULTS_XLSX),
            "external_v1": str(EXTERNAL_V1_XLSX),
            "kernel_v2": str(KERNEL_V2_XLSX),
        },
        "validation": validation_payload,
        "copied_figures": [str(path) for path in copied_figures],
        "top3_by_blind_test_rmse": combined_df.head(3)[["label", "target_test_pooled_rmse"]].to_dict(orient="records"),
    }
    JSON_OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    ensure_dirs()
    copied_figures = copy_figures()
    combined_df, internal_df, external_df, kernel_df = make_combined_summary()
    validation_payload = validate(combined_df, internal_df)
    write_excel(combined_df, internal_df, external_df, kernel_df, validation_payload)
    build_doc(combined_df, internal_df, external_df, kernel_df, copied_figures)
    write_json_summary(combined_df, validation_payload, copied_figures)


if __name__ == "__main__":
    main()
