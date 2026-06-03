from pathlib import Path
import json

import numpy as np
import pandas as pd
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"C:\Users\ROG\Desktop\paper")
CODEX_DIR = ROOT / "codex"
PRED_CSV = CODEX_DIR / "paper-update-20260420" / "target_test_predictions_all_models_20260420.csv"
CI_XLSX = CODEX_DIR / "visualization-update-20260420" / "teal_ci_visualization_summary_20260420.xlsx"
OUT_DIR = CODEX_DIR / "paper-draft-v2-20260420"
FIG_DIR = OUT_DIR / "figures"
OUT_DOCX = OUT_DIR / "method_experiment_draft_v2_20260420.docx"
OUT_XLSX = OUT_DIR / "manuscript_v2_metrics_summary_20260420.xlsx"
OUT_JSON = OUT_DIR / "manuscript_v2_generation_summary_20260420.json"

MODELS = [
    "final-retained-v1",
    "no-physical-base-v1",
    "no-mmd-v1",
    "no-attention-v1",
]

MODEL_LABELS = {
    "final-retained-v1": "Final retained model",
    "no-physical-base-v1": "w/o physical base",
    "no-mmd-v1": "w/o MMD",
    "no-attention-v1": "w/o Attention",
}

COLORS = {
    "true": "#173F3A",
    "pred": "#009C9A",
    "band": "#C7EFE9",
    "err": "#2E817C",
    "grid": "#CFE5E1",
    "text": "#102D2B",
    "final-retained-v1": "#00796B",
    "no-physical-base-v1": "#D98B3A",
    "no-mmd-v1": "#4F83A8",
    "no-attention-v1": "#7FA66A",
}

DISPLAY_SMOOTH = True
SAVGOL_WINDOW = 71
SAVGOL_POLY = 3
ERROR_STRIDE_MIN = 14
ERROR_MAX_BARS = 70

plt.rcParams["font.family"] = "Times New Roman"
plt.rcParams["font.size"] = 12
plt.rcParams["axes.linewidth"] = 1.15
plt.rcParams["figure.dpi"] = 300
plt.rcParams["savefig.dpi"] = 300
plt.rcParams["axes.unicode_minus"] = False


def ensure_dirs():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)


def rmse(y_true, y_pred):
    y_true = np.asarray(y_true).reshape(-1)
    y_pred = np.asarray(y_pred).reshape(-1)
    return float(np.sqrt(np.mean((y_pred - y_true) ** 2)))


def mae(y_true, y_pred):
    y_true = np.asarray(y_true).reshape(-1)
    y_pred = np.asarray(y_pred).reshape(-1)
    return float(np.mean(np.abs(y_pred - y_true)))


def mape(y_true, y_pred):
    y_true = np.asarray(y_true).reshape(-1)
    y_pred = np.asarray(y_pred).reshape(-1)
    denom = np.maximum(np.abs(y_true), 1e-8)
    return float(np.mean(np.abs((y_pred - y_true) / denom)) * 100.0)


def smooth_line(y, window=SAVGOL_WINDOW, poly=SAVGOL_POLY):
    y = np.asarray(y).reshape(-1)
    if not DISPLAY_SMOOTH or len(y) < 9:
        return y.copy()
    try:
        from scipy.signal import savgol_filter

        win = min(window, len(y) if len(y) % 2 == 1 else len(y) - 1)
        if win <= poly:
            win = poly + 3 if (poly + 3) % 2 == 1 else poly + 4
        win = min(win, len(y) if len(y) % 2 == 1 else len(y) - 1)
        if win < 7:
            return y.copy()
        return savgol_filter(y, window_length=win, polyorder=min(poly, win - 2), mode="interp")
    except Exception:
        win = min(window, len(y))
        if win < 5:
            return y.copy()
        return pd.Series(y).rolling(window=win, center=True, min_periods=1).mean().to_numpy()


def load_ci_half_widths():
    fallback = {
        "final-retained-v1": 0.0140,
        "no-physical-base-v1": 0.0304,
        "no-mmd-v1": 0.0139,
        "no-attention-v1": 0.0177,
    }
    if not CI_XLSX.exists():
        return fallback
    df = pd.read_excel(CI_XLSX, sheet_name="overall")
    out = fallback.copy()
    for _, row in df.iterrows():
        model = row["model"]
        if model in out:
            out[model] = float(row["ci_half_width_95"])
    return out


def compute_cell_metrics(df, model):
    rows = []
    for battery, g in df.groupby("battery_index"):
        yt = g["y_true"].to_numpy(dtype=float)
        yp = g[model].to_numpy(dtype=float)
        abs_err = np.abs(yp - yt)
        rows.append(
            {
                "model": model,
                "battery_index": int(battery),
                "n_points": int(len(g)),
                "rmse": rmse(yt, yp),
                "mae": mae(yt, yp),
                "max_abs_error": float(np.max(abs_err)),
            }
        )
    return pd.DataFrame(rows)


def compute_overall_metrics(df):
    rows = []
    cell_tables = {}
    for model in MODELS:
        cells = compute_cell_metrics(df, model)
        cell_tables[model] = cells
        rows.append(
            {
                "model": model,
                "label": MODEL_LABELS[model],
                "pooled_rmse": rmse(df["y_true"].to_numpy(), df[model].to_numpy()),
                "macro_rmse": float(cells["rmse"].mean()),
                "mae": mae(df["y_true"].to_numpy(), df[model].to_numpy()),
                "mape_percent": mape(df["y_true"].to_numpy(), df[model].to_numpy()),
            }
        )
    return pd.DataFrame(rows), cell_tables


def plot_model_2x4(df, model, ci_half, out_path):
    batteries = sorted(df["battery_index"].unique().tolist())
    fig, axes = plt.subplots(2, 4, figsize=(20, 9.2), constrained_layout=False)
    axes = axes.flatten()

    legend_handles = [
        Line2D([0], [0], color=COLORS["true"], lw=2.0, label="Measured SOH"),
        Line2D([0], [0], color=COLORS["pred"], lw=2.0, label="Predicted SOH (display-smoothed)"),
        Patch(facecolor=COLORS["band"], edgecolor=COLORS["band"], alpha=0.38, label="95% empirical band"),
        Patch(facecolor=COLORS["err"], edgecolor=COLORS["err"], alpha=0.16, label="Absolute error bars"),
    ]

    all_abs_err = np.abs(df[model].to_numpy(dtype=float) - df["y_true"].to_numpy(dtype=float))
    err_ylim = max(0.08, float(np.percentile(all_abs_err, 99)) * 3.0)

    cell_rows = []
    for i, battery in enumerate(batteries):
        ax = axes[i]
        ax_err = ax.twinx()
        g = df[df["battery_index"] == battery].sort_values("cycle")
        x = g["cycle"].to_numpy(dtype=float)
        yt = g["y_true"].to_numpy(dtype=float)
        yp_raw = g[model].to_numpy(dtype=float)
        yp_show = smooth_line(yp_raw)
        lo_show = smooth_line(yp_raw - ci_half)
        hi_show = smooth_line(yp_raw + ci_half)
        abs_err = np.abs(yp_raw - yt)

        stride = max(ERROR_STRIDE_MIN, len(x) // ERROR_MAX_BARS)
        xs = x[::stride]
        es = abs_err[::stride]
        width = max(5.0, float(np.median(np.diff(x))) * stride * 0.75) if len(x) > 1 else 5.0

        ax_err.bar(xs, es, width=width, color=COLORS["err"], alpha=0.16, linewidth=0, zorder=0)
        ax_err.set_ylim(0, err_ylim)
        ax_err.set_ylabel("Abs. Error", fontsize=9, color=COLORS["err"])
        ax_err.tick_params(axis="y", labelsize=8, colors=COLORS["err"])
        ax_err.grid(False)
        ax_err.set_zorder(0)
        ax.set_zorder(1)
        ax.patch.set_alpha(0)

        ax.fill_between(x, lo_show, hi_show, color=COLORS["band"], alpha=0.38, linewidth=0, zorder=1)
        ax.plot(x, yt, color=COLORS["true"], linewidth=2.0, zorder=3)
        ax.plot(x, yp_show, color=COLORS["pred"], linewidth=2.0, zorder=4)

        cell_rmse = rmse(yt, yp_raw)
        cell_mae = mae(yt, yp_raw)
        cell_rows.append({"battery_index": int(battery), "rmse": cell_rmse, "mae": cell_mae, "n_points": int(len(g))})

        y_min = min(float(np.min(yt)), float(np.min(lo_show)), float(np.min(yp_show)))
        y_max = max(float(np.max(yt)), float(np.max(hi_show)), float(np.max(yp_show)))
        pad = max(0.025, 0.06 * (y_max - y_min))
        ax.set_ylim(y_min - pad, y_max + pad)
        ax.set_title(f"Cell {int(battery)} | RMSE = {cell_rmse:.4f}", fontsize=13, fontweight="bold", color=COLORS["text"])
        ax.set_xlabel("Cycle", fontsize=11)
        ax.set_ylabel("SOH", fontsize=11)
        ax.grid(True, linestyle="--", color=COLORS["grid"], alpha=0.55)
        ax.tick_params(axis="both", labelsize=10)

    for j in range(len(batteries), 8):
        axes[j].axis("off")

    pooled = rmse(df["y_true"].to_numpy(), df[model].to_numpy())
    macro = float(np.mean([r["rmse"] for r in cell_rows]))
    mae_all = mae(df["y_true"].to_numpy(), df[model].to_numpy())

    fig.legend(handles=legend_handles, loc="upper center", bbox_to_anchor=(0.5, 0.985), ncol=4, frameon=False, fontsize=12)
    fig.suptitle(
        f"Batch 10 SOH Prediction - {MODEL_LABELS[model]}\nClean 2 x 4 visualization with empirical band and absolute error bars",
        fontsize=20,
        y=1.035,
        fontweight="bold",
        color=COLORS["text"],
    )
    fig.text(
        0.5,
        0.018,
        f"Pooled RMSE = {pooled:.5f}   |   Macro RMSE = {macro:.5f}   |   MAE = {mae_all:.5f}   |   95% half-width = {ci_half:.5f}",
        ha="center",
        fontsize=13,
        color=COLORS["text"],
    )
    plt.tight_layout(rect=[0, 0.055, 1, 0.925])
    fig.savefig(out_path, bbox_inches="tight")
    plt.close(fig)
    return pd.DataFrame(cell_rows)


def plot_per_battery_bar(rmse_wide, out_path):
    x = np.arange(len(rmse_wide["battery_index"]))
    width = 0.2
    fig, ax = plt.subplots(figsize=(13.5, 5.8))
    for idx, model in enumerate(MODELS):
        ax.bar(
            x + (idx - 1.5) * width,
            rmse_wide[model],
            width=width,
            label=MODEL_LABELS[model],
            color=COLORS[model],
            alpha=0.88,
            edgecolor="white",
            linewidth=0.6,
        )
    ax.set_xticks(x)
    ax.set_xticklabels(rmse_wide["battery_index"].astype(str).tolist())
    ax.set_xlabel("Battery index in Batch 10", fontsize=12)
    ax.set_ylabel("Per-battery RMSE", fontsize=12)
    ax.set_title("Per-battery blind-test RMSE comparison under strict Scheme A", fontsize=15, color=COLORS["text"])
    ax.grid(axis="y", linestyle="--", color=COLORS["grid"], alpha=0.65)
    ax.legend(frameon=False, ncol=2, fontsize=10)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    fig.savefig(out_path, dpi=300)
    plt.close(fig)


def set_font(run, name="Microsoft YaHei", size=10.5, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    r = p.add_run(text)
    set_font(r, size=14 if level == 1 else 12, bold=True)


def add_para(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_font(r)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    r = p.add_run(text)
    set_font(r, name="Consolas", size=9)


def fill_table(table, rows):
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(str(value))
            set_font(r, size=9, bold=(i == 0))


def write_docx(overall_df, rmse_wide, fig_paths, bar_path):
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("方法部分与实验设置初稿 v2")
    set_font(tr, size=16, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("基于严格 Scheme A 协议、消融实验与更新版可视化结果")
    set_font(sr, size=10)

    add_heading(doc, "1 Methodology", level=1)
    add_heading(doc, "1.1 Overall Framework", level=2)
    add_para(
        doc,
        "本文构建一种面向跨批次锂离子电池 SOH 估计的物理引导混合迁移学习框架。整体方法由全局物理退化基座、残差时序建模网络和生命周期自适应域对齐约束组成。三者承担不同层级的功能：物理基座刻画跨电池共享的长期退化趋势，残差时序网络学习围绕该基准线的个体化局部偏差，MMD 对齐项在训练阶段缓解 source domain 与 target-domain development set 之间的特征分布差异。",
    )
    add_para(
        doc,
        "设输入序列片段为 x_i，归一化循环位置为 tau_i，对应 SOH 标签为 y_i。模型的输出由物理基准预测和残差修正项相加得到：",
    )
    add_code(doc, r"\hat{y}_i = f_{phy}(\tau_i) + f_{res}(x_i,\tau_i)")
    add_para(
        doc,
        "该分解使神经网络不必独立承担完整 SOH 曲线拟合，而是集中学习难以由低阶退化函数表达的局部动态波动。消融实验中，移除物理基座会使 Batch 10 blind-test RMSE 从 0.009499 上升至 0.020716，说明宏观退化基准线是当前方法中最关键的结构支撑。",
    )

    add_heading(doc, "1.2 Global Physical Degradation Base", level=2)
    add_para(
        doc,
        "电池容量衰减具有明显的长期趋势性。为避免纯数据驱动网络将宏观退化、批次差异和局部噪声混合到同一隐空间中，本文首先使用低复杂度全局物理退化基座为所有样本建立共享参考曲线。当前实现采用三次多项式形式：",
    )
    add_code(doc, r"f_{phy}(\tau)=a\tau^3+b\tau^2+c\tau+d")
    add_para(
        doc,
        "其中 tau 为归一化循环位置，a、b、c、d 为由 source domain 数据估计得到的全局系数。该三次多项式并不被视为复杂原创结构，而是在当前任务中承担明确功能的物理退化基准线。它提供宏观趋势先验，使残差网络主要处理个体化偏差。",
    )

    add_heading(doc, "1.3 Residual Temporal Modeling", level=2)
    add_para(
        doc,
        "在物理基座给出长期退化趋势后，仍需对不同电池在局部循环阶段表现出的个体化波动进行建模。本文采用由 CNN、LSTM 和 Attention 组成的残差时序网络。CNN 提取局部序列模式，LSTM 建模窗口内时序依赖，Attention 对不同时刻隐状态进行加权聚合，以形成更集中的残差表征。",
    )
    add_para(
        doc,
        "Attention 在本文中定位为辅助表征模块，而不是核心原创机制。严格协议下的消融结果显示，移除 Attention 后 Batch 10 pooled RMSE 上升至 0.011677，说明其对残差特征聚合具有明确贡献。因此论文中应将其表述为显著有效的辅助时序表征组件，而不应夸大为独立的重大原创架构。",
    )

    add_heading(doc, "1.4 Adaptive Domain Alignment", level=2)
    add_para(
        doc,
        "跨批次 SOH 估计面临 source batches 与目标批次之间的分布偏移。为提高目标域泛化能力，本文在训练阶段引入生命周期自适应 MMD 对齐项，使 source domain 与 target-domain development set 的中间特征分布保持适度接近。训练目标可概括为：",
    )
    add_code(doc, r"\mathcal{L}=\mathcal{L}_{MSE}^{source}+\lambda(\bar{s}_t)\mathcal{D}_{MMD}^{2}(z_s,z_t)")
    add_para(
        doc,
        "其中 z_s 和 z_t 分别表示 source 与 target-domain development samples 的中间表征，lambda 表示随目标域退化阶段变化的自适应权重。需要强调的是，MMD 只在训练阶段使用 Batch 9 target-domain development set 进行对齐；Batch 10 始终保留为最终 blind-test，不参与训练、调参或 checkpoint selection。当前消融结果显示，移除 MMD 后 Batch 10 pooled RMSE 为 0.009677，较完整模型小幅变差，因此 MMD 应被描述为轻量但正向的跨域增强模块。",
    )

    add_heading(doc, "1.5 Optimization and Model Selection", level=2)
    add_para(
        doc,
        "训练过程遵循严格 Scheme A 协议。模型首先在 source domain 上进行监督回归学习，随后结合 target-domain development set 的特征对齐约束，并根据该 development set 上的 RMSE 选择最佳 checkpoint。最终性能仅在 Batch 10 blind-test set 上评估一次。该协议避免了使用测试集选择 checkpoint 的信息泄露风险，更适合作为论文主结果的实验依据。当前 smoothness penalty 不纳入最终保留模型。",
    )

    add_heading(doc, "2 Experimental Settings", level=1)
    add_heading(doc, "2.1 Dataset and Splitting Protocol", level=2)
    add_para(
        doc,
        "实验采用商业 LFP/graphite 锂离子电池寿命数据集，当前预处理版本包含 77 个电池单体。固定划分协议如下：Batch 1-8 作为 source domain，用于监督训练；Batch 9 作为 target-domain development set，用于目标域自适应和 checkpoint selection；Batch 10 作为 target-domain blind-test set，仅用于最终泛化评估。",
    )
    add_para(
        doc,
        "当前数据规模为 source domain 115256 个样本、61 个电池；target-domain development set 16310 个样本、8 个电池；target-domain blind-test set 14556 个样本、8 个电池。输入窗口长度为 100，标签采用 nominal SOH，循环位置归一化后输入物理基座和位置表征分支。",
    )

    add_heading(doc, "2.2 Compared Models", level=2)
    add_bullet(doc, "Final retained model：包含三次物理退化基座、残差时序网络、Attention 和生命周期自适应 MMD。")
    add_bullet(doc, "w/o physical base：移除三次物理退化基座，由残差网络直接预测 SOH。")
    add_bullet(doc, "w/o MMD：关闭训练阶段 MMD 对齐，保留物理基座和 Attention。")
    add_bullet(doc, "w/o Attention：移除 Attention，使用 LSTM 最后时刻隐状态作为时序表征。")

    add_heading(doc, "2.3 Training and Evaluation Metrics", level=2)
    add_para(
        doc,
        "所有模型在相同数据划分下训练 150 个 epoch，batch size 为 128，初始学习率为 1e-3，并在后期进行学习率衰减。最佳模型由 target-domain development set RMSE 选择。评价指标包括 pooled RMSE、macro RMSE、MAE、MAPE 和 per-battery RMSE。pooled RMSE 反映所有测试样本的整体误差，macro RMSE 先计算每个电池单体 RMSE 后再平均，可降低样本数差异造成的偏置。",
    )
    rows = [["Model", "Pooled RMSE", "Macro RMSE", "MAE", "MAPE (%)"]]
    for _, row in overall_df.iterrows():
        rows.append(
            [
                row["label"],
                f"{row['pooled_rmse']:.6f}",
                f"{row['macro_rmse']:.6f}",
                f"{row['mae']:.6f}",
                f"{row['mape_percent']:.3f}",
            ]
        )
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    fill_table(table, rows)

    add_heading(doc, "2.4 Visualization for Blind-test Prediction", level=2)
    add_para(
        doc,
        "为展示模型在 Batch 10 不同电池单体上的预测行为，本文绘制 2 x 4 SOH 轨迹图。每个子图包含 measured SOH、显示层面平滑后的 predicted SOH、基于 Batch 9 residual standard deviation 构造的 95% empirical band，以及非负绝对误差矩形条。需要说明的是，平滑仅用于图像展示，所有 RMSE、MAE 和 MAPE 均基于原始预测值计算。",
    )
    doc.add_picture(str(fig_paths["final-retained-v1"]), width=Inches(6.7))
    add_para(doc, "Figure 1. Batch 10 blind-test SOH prediction of the final retained model.")
    doc.add_picture(str(bar_path), width=Inches(6.7))
    add_para(doc, "Figure 2. Per-battery RMSE comparison among the retained model and ablation variants.")
    for fig_idx, model in enumerate(["no-physical-base-v1", "no-mmd-v1", "no-attention-v1"], start=3):
        doc.add_picture(str(fig_paths[model]), width=Inches(6.7))
        add_para(doc, f"Figure {fig_idx}. Batch 10 blind-test SOH prediction of {MODEL_LABELS[model]}.")
    add_para(
        doc,
        "从 per-battery RMSE 对比可以看出，移除物理基座在几乎所有目标测试电池上均造成明显误差上升，说明该模块对跨批次盲测稳定性具有关键作用。移除 Attention 后误差也出现可观增加，而移除 MMD 的影响相对较小但总体仍为负向变化。因此，当前实验结果支持将物理基座作为方法主线，将 Attention 描述为显著有效的辅助表征模块，将 MMD 描述为轻量正向的跨域增强模块。",
    )

    doc.save(OUT_DOCX)


def main():
    ensure_dirs()
    if not PRED_CSV.exists():
        raise FileNotFoundError(f"Missing prediction CSV: {PRED_CSV}")

    df = pd.read_csv(PRED_CSV)
    required = {"battery_index", "cycle", "y_true", *MODELS}
    missing = required.difference(df.columns)
    if missing:
        raise ValueError(f"Missing columns in prediction CSV: {sorted(missing)}")

    ci_widths = load_ci_half_widths()
    overall_df, cell_tables = compute_overall_metrics(df)

    fig_paths = {}
    clean_cell_tables = {}
    for model in MODELS:
        fig_path = FIG_DIR / f"{model}_batch10_2x4_clean_abs_error_bars_v2.png"
        clean_cell_tables[model] = plot_model_2x4(df, model, ci_widths[model], fig_path)
        fig_paths[model] = fig_path

    rmse_wide = pd.DataFrame({"battery_index": sorted(df["battery_index"].unique().astype(int).tolist())})
    rmse_wide["n_points"] = rmse_wide["battery_index"].map(
        cell_tables["final-retained-v1"].set_index("battery_index")["n_points"].to_dict()
    )
    for model in MODELS:
        values = cell_tables[model].set_index("battery_index")["rmse"].to_dict()
        rmse_wide[model] = rmse_wide["battery_index"].map(values)

    bar_path = FIG_DIR / "per_battery_rmse_comparison_bar_v2.png"
    plot_per_battery_bar(rmse_wide, bar_path)

    delta_df = rmse_wide[["battery_index", "n_points"]].copy()
    for model in ["no-physical-base-v1", "no-mmd-v1", "no-attention-v1"]:
        delta_df[f"{model}_minus_final"] = rmse_wide[model] - rmse_wide["final-retained-v1"]

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        overall_df.to_excel(writer, sheet_name="overall_metrics", index=False)
        rmse_wide.to_excel(writer, sheet_name="per_battery_rmse", index=False)
        delta_df.to_excel(writer, sheet_name="delta_vs_final", index=False)
        for model, table in cell_tables.items():
            table.to_excel(writer, sheet_name=model[:31], index=False)

    write_docx(overall_df, rmse_wide, fig_paths, bar_path)

    summary = {
        "status": "success",
        "output_dir": str(OUT_DIR),
        "docx": str(OUT_DOCX),
        "xlsx": str(OUT_XLSX),
        "figures": {model: str(path) for model, path in fig_paths.items()},
        "bar_figure": str(bar_path),
        "overall_metrics": overall_df.to_dict(orient="records"),
        "note": "Metrics use raw predictions. 2x4 prediction lines and empirical bands are display-smoothed only.",
    }
    OUT_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
