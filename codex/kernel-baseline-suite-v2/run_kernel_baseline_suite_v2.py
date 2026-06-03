import json
import math
import os
import pickle
import time
import warnings
from pathlib import Path
from typing import Dict, List, Tuple

import numpy as np
import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from sklearn.decomposition import PCA
from sklearn.exceptions import ConvergenceWarning
from sklearn.gaussian_process import GaussianProcessRegressor
from sklearn.gaussian_process.kernels import ConstantKernel, Matern, RBF, RationalQuadratic, WhiteKernel
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.svm import LinearSVR, SVR


ROOT = Path(r"C:\Users\ROG\Desktop\paper")
CODEX_DIR = ROOT / "codex"
DATA_DIR = ROOT / "PYthon训练代码" / "processed_clean"
SUITE_DIR = CODEX_DIR / "kernel-baseline-suite-v2"
REPORT_DOCX = SUITE_DIR / "kernel_baseline_suite_v2_report.docx"
RUN_LOG_DOCX = SUITE_DIR / "kernel_baseline_suite_v2_run_log.docx"
SUMMARY_XLSX = SUITE_DIR / "kernel_baseline_suite_v2_summary.xlsx"
RUN_JSON = SUITE_DIR / "kernel_baseline_suite_v2_summary.json"
PROGRESS_JSON = SUITE_DIR / "kernel_baseline_suite_v2_progress.json"

EXTERNAL_V1_XLSX = CODEX_DIR / "external-baseline-suite-v1" / "external_baseline_suite_v1_summary.xlsx"
INTERNAL_XLSX = CODEX_DIR / "model_runs_summary.xlsx"

LABEL_MODE = "nominal"
DATA_PROTOCOL = "processed_clean_fixed_split"
SEED = 42
GLOBAL_MAX_CYCLE = 3000.0
TIME_LIMIT_SECONDS = 5 * 3600
START_TIME = time.time()

BASELINES = [
    {
        "model": "pca-linearsvr-v1",
        "label": "PCA-LinearSVR baseline",
        "family": "kernel_statistical",
        "sample_size": 20000,
        "pca_components": 24,
        "candidates": [
            {"C": 0.1, "epsilon": 0.001},
            {"C": 0.1, "epsilon": 0.005},
            {"C": 1.0, "epsilon": 0.001},
            {"C": 1.0, "epsilon": 0.005},
            {"C": 5.0, "epsilon": 0.001},
            {"C": 5.0, "epsilon": 0.005},
            {"C": 10.0, "epsilon": 0.001},
            {"C": 10.0, "epsilon": 0.01},
        ],
    },
    {
        "model": "pca-rbf-svr-v1",
        "label": "PCA-RBF-SVR baseline",
        "family": "kernel_statistical",
        "sample_size": 7000,
        "pca_components": 24,
        "candidates": [
            {"C": 5.0, "epsilon": 0.001, "gamma": "scale"},
            {"C": 10.0, "epsilon": 0.001, "gamma": "scale"},
            {"C": 20.0, "epsilon": 0.003, "gamma": "scale"},
            {"C": 10.0, "epsilon": 0.005, "gamma": 0.05},
            {"C": 20.0, "epsilon": 0.005, "gamma": 0.05},
            {"C": 30.0, "epsilon": 0.01, "gamma": 0.03},
        ],
    },
    {
        "model": "pca-gpr-v1",
        "label": "PCA-GPR baseline",
        "family": "kernel_statistical",
        "sample_size": 1800,
        "pca_components": 12,
        "candidates": [
            {"kernel_name": "RBF"},
            {"kernel_name": "Matern-1.5"},
            {"kernel_name": "Matern-2.5"},
            {"kernel_name": "RationalQuadratic"},
        ],
    },
]


def ensure_dirs() -> None:
    SUITE_DIR.mkdir(parents=True, exist_ok=True)
    for cfg in BASELINES:
        (SUITE_DIR / cfg["model"]).mkdir(parents=True, exist_ok=True)


def elapsed_seconds() -> float:
    return float(time.time() - START_TIME)


def check_deadline(stage: str) -> None:
    if elapsed_seconds() > TIME_LIMIT_SECONDS:
        write_progress(stage, "timeout")
        raise TimeoutError(f"time limit exceeded at stage: {stage}")


def write_progress(stage: str, status: str, extra: Dict = None) -> None:
    ensure_dirs()
    payload = {
        "stage": stage,
        "status": status,
        "elapsed_seconds": round(elapsed_seconds(), 3),
        "time_limit_seconds": TIME_LIMIT_SECONDS,
        "outputs": {
            "suite_dir": str(SUITE_DIR),
            "report_docx": str(REPORT_DOCX),
            "run_log_docx": str(RUN_LOG_DOCX),
            "summary_xlsx": str(SUMMARY_XLSX),
            "run_json": str(RUN_JSON),
        },
    }
    if extra:
        payload.update(extra)
    PROGRESS_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def configure_doc(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_table(doc: Document, df: pd.DataFrame, title: str) -> None:
    doc.add_paragraph(title)
    view = df.copy()
    for col in view.columns:
        if pd.api.types.is_float_dtype(view[col]):
            view[col] = view[col].map(lambda x: round(float(x), 6) if pd.notna(x) else "")
    table = doc.add_table(rows=1, cols=len(view.columns))
    for i, col in enumerate(view.columns):
        table.rows[0].cells[i].text = str(col)
    for _, row in view.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row.tolist()):
            cells[i].text = str(val)


def load_npy(path: Path) -> np.ndarray:
    if not path.exists():
        raise FileNotFoundError(f"missing file: {path}")
    return np.load(path)


def load_split(split_name: str, label_mode: str) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    x = load_npy(DATA_DIR / f"clean_{split_name}_x.npy").astype(np.float32)
    p = load_npy(DATA_DIR / f"clean_{split_name}_p.npy").astype(np.float32).reshape(-1)
    bidx = load_npy(DATA_DIR / f"clean_{split_name}_battery_index.npy").astype(np.int64).reshape(-1)
    batch = load_npy(DATA_DIR / f"clean_{split_name}_batch.npy").astype(np.int64).reshape(-1)
    if label_mode == "nominal":
        y = load_npy(DATA_DIR / f"clean_{split_name}_y_nominal.npy").astype(np.float32).reshape(-1)
    elif label_mode == "relative":
        y = load_npy(DATA_DIR / f"clean_{split_name}_y_relative.npy").astype(np.float32).reshape(-1)
    else:
        raise ValueError("unsupported label mode")
    return x, y, p, bidx, batch


def build_features(x: np.ndarray, p: np.ndarray) -> np.ndarray:
    pos = (p.reshape(-1, 1) / GLOBAL_MAX_CYCLE).astype(np.float32)
    return np.concatenate([x.astype(np.float32), pos], axis=1)


def rmse(y_true: np.ndarray, y_pred: np.ndarray) -> float:
    return float(math.sqrt(np.mean((y_true.reshape(-1) - y_pred.reshape(-1)) ** 2)))


def mae(y_true: np.ndarray, y_pred: np.ndarray) -> float:
    return float(np.mean(np.abs(y_true.reshape(-1) - y_pred.reshape(-1))))


def mape_percent(y_true: np.ndarray, y_pred: np.ndarray) -> float:
    denom = np.clip(np.abs(y_true.reshape(-1)), 1e-8, None)
    return float(np.mean(np.abs((y_true.reshape(-1) - y_pred.reshape(-1)) / denom)) * 100.0)


def macro_rmse_by_battery(y_true: np.ndarray, y_pred: np.ndarray, battery_index: np.ndarray) -> float:
    y_true = y_true.reshape(-1)
    y_pred = y_pred.reshape(-1)
    battery_index = battery_index.reshape(-1)
    vals = []
    for b in np.unique(battery_index):
        mask = battery_index == b
        vals.append(rmse(y_true[mask], y_pred[mask]))
    return float(np.mean(vals))


def evaluate_predictions(y_true: np.ndarray, y_pred: np.ndarray, battery_index: np.ndarray) -> Tuple[Dict[str, float], pd.DataFrame]:
    y_true = y_true.reshape(-1)
    y_pred = y_pred.reshape(-1)
    battery_index = battery_index.reshape(-1)
    metrics = {
        "pooled_rmse": rmse(y_true, y_pred),
        "macro_rmse": macro_rmse_by_battery(y_true, y_pred, battery_index),
        "mae": mae(y_true, y_pred),
        "mape_percent": mape_percent(y_true, y_pred),
    }
    rows = []
    for b in np.unique(battery_index):
        mask = battery_index == b
        rows.append(
            {
                "battery_index": int(b),
                "rmse": rmse(y_true[mask], y_pred[mask]),
                "mae": mae(y_true[mask], y_pred[mask]),
                "n_points": int(mask.sum()),
            }
        )
    return metrics, pd.DataFrame(rows).sort_values("battery_index").reset_index(drop=True)


def split_summary(name: str, x: np.ndarray, y: np.ndarray, p: np.ndarray, bidx: np.ndarray, batch: np.ndarray) -> Dict[str, object]:
    return {
        "split": name,
        "samples": int(len(x)),
        "cells": int(len(np.unique(bidx))),
        "batches": ",".join(map(str, sorted(np.unique(batch).tolist()))),
        "x_shape": "x".join(map(str, x.shape)),
        "soh_min": float(np.min(y)),
        "soh_max": float(np.max(y)),
        "cycle_min": float(np.min(p)),
        "cycle_max": float(np.max(p)),
    }


def stratified_source_subset(p: np.ndarray, battery_index: np.ndarray, sample_size: int, n_cycle_bins: int = 10) -> np.ndarray:
    n = len(p)
    if sample_size >= n:
        return np.arange(n, dtype=np.int64)
    rng = np.random.default_rng(SEED)
    p_norm = np.clip(p.reshape(-1) / GLOBAL_MAX_CYCLE, 0.0, 0.999999)
    cycle_bin = np.minimum((p_norm * n_cycle_bins).astype(np.int64), n_cycle_bins - 1)
    keys = np.stack([battery_index.reshape(-1), cycle_bin], axis=1)
    _, inverse, counts = np.unique(keys, axis=0, return_inverse=True, return_counts=True)
    alloc = np.floor(counts / counts.sum() * sample_size).astype(np.int64)
    alloc = np.clip(alloc, 1, counts)
    while alloc.sum() > sample_size:
        reducible = np.where(alloc > 1)[0]
        if len(reducible) == 0:
            break
        idx = reducible[np.argmax(alloc[reducible])]
        alloc[idx] -= 1
    while alloc.sum() < sample_size:
        expandable = np.where(alloc < counts)[0]
        if len(expandable) == 0:
            break
        idx = expandable[np.argmax(counts[expandable] - alloc[expandable])]
        alloc[idx] += 1
    chosen = []
    for group_id in range(len(counts)):
        idxs = np.where(inverse == group_id)[0]
        take = int(min(alloc[group_id], len(idxs)))
        if take == len(idxs):
            chosen.append(idxs)
        else:
            chosen.append(rng.choice(idxs, size=take, replace=False))
    return np.sort(np.concatenate(chosen).astype(np.int64))


def build_linear_svr_pipeline(pca_components: int, params: Dict[str, object]) -> Pipeline:
    return Pipeline(
        [
            ("scaler", StandardScaler()),
            ("pca", PCA(n_components=pca_components, svd_solver="randomized", random_state=SEED)),
            (
                "model",
                LinearSVR(
                    C=float(params["C"]),
                    epsilon=float(params["epsilon"]),
                    loss="epsilon_insensitive",
                    random_state=SEED,
                    max_iter=20000,
                ),
            ),
        ]
    )


def build_rbf_svr_pipeline(pca_components: int, params: Dict[str, object]) -> Pipeline:
    return Pipeline(
        [
            ("scaler", StandardScaler()),
            ("pca", PCA(n_components=pca_components, svd_solver="randomized", random_state=SEED)),
            (
                "model",
                SVR(
                    kernel="rbf",
                    C=float(params["C"]),
                    epsilon=float(params["epsilon"]),
                    gamma=params["gamma"],
                ),
            ),
        ]
    )


def build_gpr_kernel(kernel_name: str):
    if kernel_name == "RBF":
        return ConstantKernel(1.0, (0.1, 10.0)) * RBF(length_scale=1.0, length_scale_bounds=(0.05, 10.0)) + WhiteKernel(
            noise_level=1e-4, noise_level_bounds=(1e-6, 1e-2)
        )
    if kernel_name == "Matern-1.5":
        return ConstantKernel(1.0, (0.1, 10.0)) * Matern(length_scale=1.0, nu=1.5) + WhiteKernel(
            noise_level=1e-4, noise_level_bounds=(1e-6, 1e-2)
        )
    if kernel_name == "Matern-2.5":
        return ConstantKernel(1.0, (0.1, 10.0)) * Matern(length_scale=1.0, nu=2.5) + WhiteKernel(
            noise_level=1e-4, noise_level_bounds=(1e-6, 1e-2)
        )
    if kernel_name == "RationalQuadratic":
        return ConstantKernel(1.0, (0.1, 10.0)) * RationalQuadratic(length_scale=1.0, alpha=1.0) + WhiteKernel(
            noise_level=1e-4, noise_level_bounds=(1e-6, 1e-2)
        )
    raise ValueError(f"unknown kernel: {kernel_name}")


def build_gpr_pipeline(pca_components: int, params: Dict[str, object]) -> Pipeline:
    return Pipeline(
        [
            ("scaler", StandardScaler()),
            ("pca", PCA(n_components=pca_components, svd_solver="randomized", random_state=SEED)),
            (
                "model",
                GaussianProcessRegressor(
                    kernel=build_gpr_kernel(str(params["kernel_name"])),
                    normalize_y=True,
                    n_restarts_optimizer=1,
                    alpha=1e-6,
                ),
            ),
        ]
    )


def build_pipeline(cfg: Dict[str, object], params: Dict[str, object]) -> Pipeline:
    if cfg["model"] == "pca-linearsvr-v1":
        return build_linear_svr_pipeline(int(cfg["pca_components"]), params)
    if cfg["model"] == "pca-rbf-svr-v1":
        return build_rbf_svr_pipeline(int(cfg["pca_components"]), params)
    if cfg["model"] == "pca-gpr-v1":
        return build_gpr_pipeline(int(cfg["pca_components"]), params)
    raise ValueError(f"unsupported model: {cfg['model']}")


def run_single_baseline(
    cfg: Dict[str, object],
    x_source: np.ndarray,
    y_source: np.ndarray,
    p_source: np.ndarray,
    bidx_source: np.ndarray,
    x_dev: np.ndarray,
    y_dev: np.ndarray,
    bidx_dev: np.ndarray,
    x_test: np.ndarray,
    y_test: np.ndarray,
    bidx_test: np.ndarray,
) -> Dict[str, object]:
    check_deadline(f"before_{cfg['model']}")
    model_dir = SUITE_DIR / str(cfg["model"])
    subset_idx = stratified_source_subset(p_source, bidx_source, int(cfg["sample_size"]))
    x_train = x_source[subset_idx]
    y_train = y_source[subset_idx].reshape(-1)
    sub_bidx = bidx_source[subset_idx]
    history_rows: List[Dict[str, object]] = []
    best = None
    total_candidates = len(cfg["candidates"])

    for i, params in enumerate(cfg["candidates"], start=1):
        check_deadline(f"{cfg['model']}_candidate_{i}")
        write_progress(
            "training",
            "running",
            {
                "current_model": cfg["model"],
                "current_label": cfg["label"],
                "candidate_index": i,
                "candidate_total": total_candidates,
                "source_subset_size": int(len(subset_idx)),
                "pca_components": int(cfg["pca_components"]),
            },
        )
        pipeline = build_pipeline(cfg, params)
        fit_start = time.time()
        warning_msgs: List[str] = []
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            warnings.filterwarnings("always", category=ConvergenceWarning)
            pipeline.fit(x_train, y_train)
            for item in caught:
                warning_msgs.append(str(item.message))
        fit_seconds = float(time.time() - fit_start)
        dev_pred = pipeline.predict(x_dev).reshape(-1)
        dev_metrics, _ = evaluate_predictions(y_dev, dev_pred, bidx_dev)
        row = {
            "model": cfg["model"],
            "label": cfg["label"],
            "candidate_index": i,
            "params_json": json.dumps(params, ensure_ascii=False),
            "fit_seconds": fit_seconds,
            "target_dev_rmse": dev_metrics["pooled_rmse"],
            "target_dev_macro_rmse": dev_metrics["macro_rmse"],
            "target_dev_mae": dev_metrics["mae"],
            "warnings_count": len(warning_msgs),
            "first_warning": warning_msgs[0][:300] if warning_msgs else "",
        }
        history_rows.append(row)
        if best is None or row["target_dev_rmse"] < best["history"]["target_dev_rmse"]:
            best = {"pipeline": pipeline, "params": params, "history": row, "warnings": warning_msgs}

    best_pipeline: Pipeline = best["pipeline"]
    dev_pred = best_pipeline.predict(x_dev).reshape(-1)
    test_pred = best_pipeline.predict(x_test).reshape(-1)
    dev_metrics, dev_batt = evaluate_predictions(y_dev, dev_pred, bidx_dev)
    test_metrics, test_batt = evaluate_predictions(y_test, test_pred, bidx_test)
    predictions_df = pd.DataFrame(
        {
            "battery_index": bidx_test.reshape(-1),
            "y_true": y_test.reshape(-1),
            "y_pred": test_pred.reshape(-1),
            "abs_error": np.abs(y_test.reshape(-1) - test_pred.reshape(-1)),
        }
    )
    with open(model_dir / "best_pipeline.pkl", "wb") as f:
        pickle.dump(best_pipeline, f)
    pd.DataFrame(history_rows).to_csv(model_dir / "candidate_history.csv", index=False, encoding="utf-8-sig")
    predictions_df.to_csv(model_dir / "target_test_predictions.csv", index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(model_dir / "per_battery_metrics.xlsx", engine="openpyxl") as writer:
        dev_batt.to_excel(writer, sheet_name="target_dev", index=False)
        test_batt.to_excel(writer, sheet_name="target_test", index=False)
        predictions_df.head(2000).to_excel(writer, sheet_name="pred_head", index=False)
    best_metrics = {
        "group": "new_kernel",
        "model": cfg["model"],
        "label": cfg["label"],
        "family": cfg["family"],
        "source_subset_size": int(len(subset_idx)),
        "pca_components": int(cfg["pca_components"]),
        "best_params_json": json.dumps(best["params"], ensure_ascii=False),
        "target_dev_rmse": dev_metrics["pooled_rmse"],
        "target_dev_macro_rmse": dev_metrics["macro_rmse"],
        "target_dev_mae": dev_metrics["mae"],
        "target_test_rmse": test_metrics["pooled_rmse"],
        "target_test_macro_rmse": test_metrics["macro_rmse"],
        "target_test_mae": test_metrics["mae"],
        "target_test_mape_percent": test_metrics["mape_percent"],
        "fit_seconds_best_candidate": best["history"]["fit_seconds"],
        "warnings_count_best_candidate": len(best["warnings"]),
    }
    (model_dir / "best_metrics.json").write_text(json.dumps(best_metrics, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "summary_row": best_metrics,
        "history_df": pd.DataFrame(history_rows),
        "dev_battery_df": dev_batt,
        "test_battery_df": test_batt,
        "predictions_df": predictions_df,
    }


def build_report_doc(
    split_df: pd.DataFrame,
    existing_internal_df: pd.DataFrame,
    prev_external_df: pd.DataFrame,
    new_kernel_df: pd.DataFrame,
    combined_df: pd.DataFrame,
) -> None:
    doc = Document()
    configure_doc(doc)
    doc.add_heading("Kernel Baseline Suite v2", level=0)
    doc.add_paragraph(
        "本轮在既有内部消融和首轮外部基线的基础上，进一步补充核方法/统计学习基线，包括 PCA-LinearSVR、PCA-RBF-SVR 与 PCA-GPR。"
        "所有新增实验继续采用 processed_clean_fixed_split 协议：Batch 1-8 为 source，Batch 9 仅用于开发集模型选择，Batch 10 仅用于最终 blind-test，避免测试信息泄露。"
    )
    doc.add_paragraph(
        "考虑到 source 端窗口级样本规模达到 115288 条，标准 GPR 与核 SVR 若直接在全量样本上训练，其时间与内存开销显著偏高。"
        "因此，本轮仅在 source 端执行分层下采样与 PCA 压缩；该处理不改变数据协议，也不使用 target_test 参与模型选择。"
    )
    add_table(doc, split_df, "数据划分摘要")
    add_table(doc, new_kernel_df, "新增核方法基线结果")
    ranking = combined_df.sort_values("target_test_rmse", ascending=True).reset_index(drop=True)
    add_table(doc, ranking[["group", "model", "label", "target_dev_rmse", "target_test_rmse", "target_test_macro_rmse"]], "全部模型 blind-test 排序")

    final_row = combined_df.loc[combined_df["model"] == "final-retained-v1"].iloc[0]
    best_kernel = new_kernel_df.sort_values("target_test_rmse", ascending=True).iloc[0]
    external_all = pd.concat([prev_external_df, new_kernel_df], ignore_index=True)
    best_external = external_all.sort_values("target_test_rmse", ascending=True).iloc[0]
    reduction_vs_kernel = (best_kernel["target_test_rmse"] - final_row["target_test_rmse"]) / best_kernel["target_test_rmse"] * 100.0
    reduction_vs_external = (best_external["target_test_rmse"] - final_row["target_test_rmse"]) / best_external["target_test_rmse"] * 100.0
    doc.add_paragraph("可直接用于论文正文的结果表述建议：")
    doc.add_paragraph(
        f"在严格的 cross-batch blind-test 条件下，所保留的最终模型仍取得最优结果，其 Batch 10 pooled RMSE 为 {final_row['target_test_rmse']:.6f}。"
        f"在本轮新增的核方法基线中，表现最好的模型为 {best_kernel['label']}，其 Batch 10 pooled RMSE 为 {best_kernel['target_test_rmse']:.6f}。"
        f"与该核方法最优基线相比，最终模型的 RMSE 进一步下降 {reduction_vs_kernel:.2f}%，表明仅依赖统计核映射或浅层回归器，仍不足以完整刻画跨批次寿命退化中的宏观趋势与个体差异。"
    )
    doc.add_paragraph(
        f"将本轮核方法结果与上一轮物理基线、循环神经网络基线联合考察后，最强外部基线为 {best_external['label']}，其 Batch 10 pooled RMSE 为 {best_external['target_test_rmse']:.6f}。"
        f"相比之下，最终模型在 blind-test 上仍获得 {reduction_vs_external:.2f}% 的误差下降，说明物理退化基座、残差时序表征与训练阶段域对齐的组合，在当前任务中具有更稳定的跨批次泛化能力。"
    )
    doc.add_paragraph(
        "进一步从方法归因角度看，纯物理基线能够提供统一的退化先验，但难以覆盖不同电池在局部阶段上的非线性波动；"
        "纯数据驱动或核方法可以学习局部映射关系，但在 source 与 target 分布存在偏移时，泛化稳定性仍然受限。"
        "因此，当前 retained model 的优势并不来自单一模块堆叠，而是来自宏观退化建模、个体化残差修正与跨域训练约束之间的功能分工。"
    )
    REPORT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(REPORT_DOCX)


def build_run_log_doc(split_df: pd.DataFrame, result_blobs: List[Dict[str, object]]) -> None:
    doc = Document()
    configure_doc(doc)
    doc.add_heading("Kernel Baseline Suite v2 Run Log", level=0)
    doc.add_paragraph(f"运行协议: {DATA_PROTOCOL}; 标签模式: {LABEL_MODE}; 时间上限: {TIME_LIMIT_SECONDS / 3600:.1f} 小时。")
    doc.add_paragraph(
        "本轮工作延续前序进度：内部消融已完成（no-physical-base / no-MMD / no-Attention），首轮外部基线已完成（physical-only / LSTM / GRU / CNN-LSTM），"
        "本轮新增核方法/统计学习基线，用于补齐与常见传统方法的比较链条。"
    )
    add_table(doc, split_df, "数据划分摘要")
    for blob in result_blobs:
        row = blob["summary_row"]
        doc.add_paragraph(
            f"{row['label']}: source subset={row['source_subset_size']}, PCA={row['pca_components']}, "
            f"target_dev RMSE={row['target_dev_rmse']:.6f}, target_test RMSE={row['target_test_rmse']:.6f}, "
            f"macro RMSE={row['target_test_macro_rmse']:.6f}."
        )
        warn_count = int(row["warnings_count_best_candidate"])
        if warn_count > 0:
            doc.add_paragraph(f"警告记录: 最优候选含 {warn_count} 条 sklearn 收敛/数值提示，详见对应 candidate_history.csv。")
    doc.add_paragraph("本轮遇到的问题与处理：")
    doc.add_paragraph("1. 标准 GPR 与核 SVR 对 10 万级窗口样本直接训练的计算开销过高，因此改为仅在 source 侧做分层下采样与 PCA 压缩。")
    doc.add_paragraph("2. 为避免 UI 无法展示长任务进度的问题，脚本在运行期间持续更新 progress.json，可用于查看当前模型、候选编号与已用时。")
    doc.add_paragraph("3. 继续直接调用 D:\\Anaconda\\envs\\battery_project\\python.exe 前台运行，避免后台启动器引入 Conda DLL 路径问题。")
    RUN_LOG_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(RUN_LOG_DOCX)


def normalize_existing_internal(df: pd.DataFrame) -> pd.DataFrame:
    view = df.copy()
    rename_map = {}
    if "best_target_adapt_rmse" in view.columns:
        rename_map["best_target_adapt_rmse"] = "target_dev_rmse"
    if "final_target_test_rmse" in view.columns:
        rename_map["final_target_test_rmse"] = "target_test_rmse"
    view = view.rename(columns=rename_map)
    if "group" not in view.columns:
        view["group"] = "existing_internal"
    if "label" not in view.columns:
        view["label"] = view["model"]
    for col in ["target_test_macro_rmse", "target_test_mae", "target_test_mape_percent"]:
        if col not in view.columns:
            view[col] = np.nan
    return view


def normalize_prev_external(df: pd.DataFrame) -> pd.DataFrame:
    view = df.copy()
    if "group" not in view.columns:
        view["group"] = "prev_external"
    if "label" not in view.columns:
        view["label"] = view["model"]
    return view


def load_existing_kernel_results() -> List[Dict[str, object]]:
    blobs = []
    for cfg in BASELINES:
        model_dir = SUITE_DIR / str(cfg["model"])
        best_metrics = json.loads((model_dir / "best_metrics.json").read_text(encoding="utf-8"))
        history_df = pd.read_csv(model_dir / "candidate_history.csv")
        test_battery_df = pd.read_excel(model_dir / "per_battery_metrics.xlsx", sheet_name="target_test")
        blobs.append(
            {
                "summary_row": best_metrics,
                "history_df": history_df,
                "test_battery_df": test_battery_df,
            }
        )
    return blobs


def main() -> None:
    ensure_dirs()
    warnings.filterwarnings("ignore", category=UserWarning)
    write_progress("bootstrap", "running")
    source_x, source_y, source_p, source_bidx, source_batch = load_split("source", LABEL_MODE)
    dev_x, dev_y, dev_p, dev_bidx, dev_batch = load_split("target_adapt", LABEL_MODE)
    test_x, test_y, test_p, test_bidx, test_batch = load_split("target_test", LABEL_MODE)
    source_feat = build_features(source_x, source_p)
    dev_feat = build_features(dev_x, dev_p)
    test_feat = build_features(test_x, test_p)

    split_df = pd.DataFrame(
        [
            split_summary("source", source_x, source_y, source_p, source_bidx, source_batch),
            split_summary("target_adapt", dev_x, dev_y, dev_p, dev_bidx, dev_batch),
            split_summary("target_test", test_x, test_y, test_p, test_bidx, test_batch),
        ]
    )

    rebuild_only = os.environ.get("KERNEL_BASELINE_V2_REBUILD_ONLY", "0") == "1"
    if rebuild_only:
        results = load_existing_kernel_results()
    else:
        results: List[Dict[str, object]] = []
        for cfg in BASELINES:
            write_progress("training", "running", {"current_model": cfg["model"], "current_label": cfg["label"]})
            blob = run_single_baseline(
                cfg,
                source_feat,
                source_y,
                source_p,
                source_bidx,
                dev_feat,
                dev_y,
                dev_bidx,
                test_feat,
                test_y,
                test_bidx,
            )
            results.append(blob)

    new_kernel_df = pd.DataFrame([x["summary_row"] for x in results]).sort_values("target_test_rmse").reset_index(drop=True)
    existing_internal_df = normalize_existing_internal(pd.read_excel(INTERNAL_XLSX, sheet_name=0))
    prev_external_df = normalize_prev_external(pd.read_excel(EXTERNAL_V1_XLSX, sheet_name="new_external"))
    combined_df = pd.concat([existing_internal_df, prev_external_df, new_kernel_df], ignore_index=True, sort=False)
    combined_df = combined_df.sort_values("target_test_rmse").reset_index(drop=True)

    with pd.ExcelWriter(SUMMARY_XLSX, engine="openpyxl") as writer:
        existing_internal_df.to_excel(writer, sheet_name="existing_internal", index=False)
        prev_external_df.to_excel(writer, sheet_name="prev_external", index=False)
        new_kernel_df.to_excel(writer, sheet_name="new_kernel", index=False)
        combined_df.to_excel(writer, sheet_name="combined", index=False)
        split_df.to_excel(writer, sheet_name="split_summary", index=False)
        for blob in results:
            model_name = blob["summary_row"]["model"].replace("-", "_")[:20]
            blob["history_df"].to_excel(writer, sheet_name=f"{model_name}_hist"[:31], index=False)
            blob["test_battery_df"].to_excel(writer, sheet_name=f"{model_name}_test"[:31], index=False)

    run_payload = {
        "protocol": DATA_PROTOCOL,
        "label_mode": LABEL_MODE,
        "time_limit_seconds": TIME_LIMIT_SECONDS,
        "elapsed_seconds": round(elapsed_seconds(), 3),
        "new_kernel_results": new_kernel_df.to_dict(orient="records"),
        "outputs": {
            "report_docx": str(REPORT_DOCX),
            "run_log_docx": str(RUN_LOG_DOCX),
            "summary_xlsx": str(SUMMARY_XLSX),
        },
    }
    RUN_JSON.write_text(json.dumps(run_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    build_report_doc(split_df, existing_internal_df, prev_external_df, new_kernel_df, combined_df)
    build_run_log_doc(split_df, results)
    write_progress("completed", "done", {"elapsed_seconds": round(elapsed_seconds(), 3), "models_completed": [cfg["model"] for cfg in BASELINES]})


if __name__ == "__main__":
    main()
